"""Excel loader (aligned with ragent ExcelDocumentParser + ExcelTableNormalizer).

Reads every VISIBLE worksheet as one block Document with kind="table". The
normalized headers/rows are carried in private metadata keys (_table_headers /
_table_rows) for the TableAwareSplitter to build the key-value vector text; the
splitter drops those private keys so they never reach the vector payload.

Formula handling uses the dual-workbook read: data_only=True gives the cached
result when present, and when the cache is missing the formula string is read
back from the data_only=False workbook (openpyxl never computes formulas).
Merged cells are expanded in the normalized matrix (top-left value copied),
never written back to MergedCell objects.
"""

from __future__ import annotations

import hashlib
import os
from datetime import date, datetime, time
from typing import Any, List

from langchain_core.documents import Document
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

from deepsearcher.loader.file_loader.base import BaseLoader
from deepsearcher.loader.file_loader.table_text import render_markdown_table

EXCEL_PARSER_VERSION = "openpyxl-table-v1"
HEADER_SEPARATOR = "|"


def _clean(value: Any) -> str:
    text = str(value or "")
    return text.replace("\x00", "").replace("\u00ad", "")


def _format_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return _clean(value).strip()


def _cell_value(values_cell: Cell, formulas_cell: Cell) -> str:
    """Cached result first, then the raw formula string, else empty."""
    value = values_cell.value
    if value is not None:
        return _format_cell(value)
    if formulas_cell.data_type == "f" and formulas_cell.value is not None:
        return str(formulas_cell.value)
    return ""


def _wrap_hyperlink(formatted: str, cell: Cell) -> str:
    target = getattr(cell.hyperlink, "target", None) if cell.hyperlink is not None else None
    if not target or not formatted:
        return formatted
    return f"[{formatted}]({target})"


def _wrap_strikethrough(formatted: str, cell: Cell) -> str:
    try:
        struck = bool(cell.font.strike)
    except Exception:
        struck = False
    if struck and formatted:
        return f"~~{formatted}~~"
    return formatted


class _NormalizedTable:
    def __init__(self, headers: List[str], rows: List[List[str]]):
        self.headers = headers
        self.rows = rows

    @property
    def is_empty(self) -> bool:
        return not self.headers and not self.rows


def _normalize_sheet(values_ws, formulas_ws, header_rows: int) -> _NormalizedTable:
    if header_rows < 1:
        raise ValueError(f"header_rows must be >= 1, got {header_rows}")
    last_row = values_ws.max_row or 0
    max_col = values_ws.max_column or 0
    if last_row < 1 or max_col < 1:
        return _NormalizedTable([], [])

    grid: list[list[str]] = []
    for r in range(1, last_row + 1):
        grid.append([])
        for c in range(1, max_col + 1):
            values_cell = values_ws.cell(row=r, column=c)
            formulas_cell = formulas_ws.cell(row=r, column=c)
            formatted = _cell_value(values_cell, formulas_cell)
            formatted = _wrap_hyperlink(formatted, values_cell)
            formatted = _wrap_strikethrough(formatted, values_cell)
            grid[-1].append(formatted)

    # Step 2: expand merged regions (top-left value copied into every cell).
    for region in values_ws.merged_cells.ranges:
        top, left = region.min_row - 1, region.min_col - 1
        value = grid[top][left]
        if not value:
            continue
        for r in range(top, min(region.max_row, last_row)):
            for c in range(left, min(region.max_col, max_col)):
                grid[r][c] = value

    # Step 3: drop fully-empty columns.
    cols = [c for c in range(max_col) if any(grid[r][c] for r in range(last_row))]
    if not cols:
        return _NormalizedTable([], [])

    # Step 4: flatten the first header_rows rows into one header line.
    effective_header_rows = min(header_rows, last_row)
    headers: List[str] = []
    for c in cols:
        parts: list[str] = []
        previous: str | None = None
        for r in range(effective_header_rows):
            value = grid[r][c]
            if not value or value == previous:
                continue
            parts.append(value)
            previous = value
        headers.append(HEADER_SEPARATOR.join(parts))

    # Step 5: collect data rows, skipping fully-empty ones.
    rows: List[List[str]] = []
    for r in range(effective_header_rows, last_row):
        row_values = [grid[r][c] for c in cols]
        if any(value for value in row_values):
            rows.append(row_values)
    return _NormalizedTable(headers, rows)


class ExcelLoader(BaseLoader):
    """xlsx loader producing one kind="table" block per visible worksheet."""

    def __init__(self, *, header_rows: int = 1):
        self.header_rows = max(1, int(header_rows))

    def load_file(self, file_path: str) -> List[Document]:
        lower_path = file_path.lower()
        if not lower_path.endswith(".xlsx"):
            raise ValueError(f"Unsupported Excel file type: {file_path}")

        with open(file_path, "rb") as source:
            document_id = hashlib.sha256(source.read()).hexdigest()

        values_wb = load_workbook(file_path, data_only=True, read_only=False)
        formulas_wb = load_workbook(file_path, data_only=False, read_only=False)
        try:
            visible_sheets = [
                sheet
                for sheet in values_wb.worksheets
                if sheet.sheet_state not in {"hidden", "veryHidden"}
            ]
            documents: List[Document] = []
            for visible_index, sheet in enumerate(visible_sheets, start=1):
                table = _normalize_sheet(
                    sheet,
                    formulas_wb[sheet.title],
                    header_rows=self.header_rows,
                )
                if table.is_empty:
                    continue
                documents.append(
                    Document(
                        page_content=render_markdown_table(table.headers, table.rows),
                        metadata={
                            "reference": file_path,
                            "document_id": document_id,
                            "display_name": os.path.basename(file_path),
                            "parser_version": EXCEL_PARSER_VERSION,
                            "extraction_method": "excel",
                            "kind": "table",
                            "sheet_name": sheet.title,
                            "sheet_index": visible_index,
                            "sheet_count": len(visible_sheets),
                            "header_rows": min(self.header_rows, (sheet.max_row or 0)),
                            "_table_headers": table.headers,
                            "_table_rows": table.rows,
                        },
                    )
                )
            return documents
        finally:
            values_wb.close()
            formulas_wb.close()

    @property
    def supported_file_types(self) -> List[str]:
        return ["xlsx"]
