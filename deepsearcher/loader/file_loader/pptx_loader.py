"""PPTX loader (block-aware: one slide may yield several block Documents).

Ordinary slide text becomes one kind="text" block; every table on the slide
becomes its own kind="table" block (headers = first row, carried in private
metadata for the TableAwareSplitter). Blocks share slide identity metadata so the
splitter never has to guess the internal structure of a slide.
"""

from __future__ import annotations

import hashlib
import os
from typing import List

from langchain_core.documents import Document

from deepsearcher.loader.file_loader.base import BaseLoader
from deepsearcher.loader.file_loader.table_text import render_markdown_table

PPTX_PARSER_VERSION = "python-pptx-v1"


def _table_rows(table) -> tuple[List[str], List[List[str]]]:
    headers: List[str] = []
    rows: List[List[str]] = []
    for row_index, row in enumerate(table.rows):
        values = [cell.text or "" for cell in row.cells]
        if row_index == 0:
            headers = values
        else:
            rows.append(values)
    return headers, rows


class PptxLoader(BaseLoader):
    """pptx loader producing slide-level text/table block Documents."""

    def load_file(self, file_path: str) -> List[Document]:
        lower_path = file_path.lower()
        if not lower_path.endswith(".pptx"):
            raise ValueError(f"Unsupported PowerPoint file type: {file_path}")

        from pptx import Presentation

        with open(file_path, "rb") as source:
            document_id = hashlib.sha256(source.read()).hexdigest()
        display_name = os.path.basename(file_path)
        presentation = Presentation(file_path)
        total_slides = len(presentation.slides)

        documents: List[Document] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            base_metadata = {
                "reference": file_path,
                "document_id": document_id,
                "display_name": display_name,
                "parser_version": PPTX_PARSER_VERSION,
                "extraction_method": "pptx",
                "slide_number": slide_number,
                "total_slides": total_slides,
            }
            text_parts: List[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_table", False) and shape.has_table:
                    headers, rows = _table_rows(shape.table)
                    if headers or rows:
                        documents.append(
                            Document(
                                page_content=render_markdown_table(headers, rows),
                                metadata={
                                    **base_metadata,
                                    "kind": "table",
                                    "_table_headers": headers,
                                    "_table_rows": rows,
                                },
                            )
                        )
                elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    text = (shape.text_frame.text or "").strip()
                    if text:
                        text_parts.append(text)
            if text_parts:
                documents.append(
                    Document(
                        page_content="\n\n".join(text_parts),
                        metadata={**base_metadata, "kind": "text"},
                    )
                )
        return documents

    @property
    def supported_file_types(self) -> List[str]:
        return ["pptx"]
