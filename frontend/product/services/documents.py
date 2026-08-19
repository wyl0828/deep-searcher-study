from __future__ import annotations

import asyncio
import hashlib
import json
import logging
import os
import secrets
import sys
import time
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from uuid import uuid4

import httpx
from fastapi import UploadFile
from sqlalchemy import func, or_, select, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from deepsearcher.loader.file_loader.mime_type import (
    EXTENSION_FAMILY,
    detect_upload_extension,
    normalize_extension,
    validate_upload,
)
from deepsearcher.versioning import normalize_version_family
from frontend.product.backend import backend_request_headers
from frontend.product.db import DATA_DIR, SessionLocal
from frontend.product.errors import ProductError
from frontend.product.messaging import MessageDispatchError, dispatch_ingest_transaction
from frontend.product.models import Document, IngestJob, KnowledgeBase, utcnow
from frontend.product.repositories import create_ingest_job
from frontend.product.services.ingestion_pipeline import (
    STATUS_FAILED,
    IngestionContext,
    NodeFailure,
    default_pipeline_steps,
    execute_chain,
)
from frontend.product.storage import LocalObjectStorage, StorageError, get_object_storage

BACKEND_URL = os.environ.get("DEEPSEARCHER_API_URL", "http://127.0.0.1:8500").rstrip("/")
logger = logging.getLogger(__name__)


def _positive_int_environment(name: str, default: int) -> int:
    try:
        value = int(os.environ.get(name, default))
    except (TypeError, ValueError):
        return default
    return value if value > 0 else default


MAX_PDF_BYTES = _positive_int_environment("DEEPSEARCHER_MAX_PDF_BYTES", 20 * 1024 * 1024)
MAX_PDF_PAGES = _positive_int_environment("DEEPSEARCHER_MAX_PDF_PAGES", 500)
MAX_DOCUMENTS_PER_KNOWLEDGE_BASE = _positive_int_environment(
    "DEEPSEARCHER_MAX_DOCUMENTS_PER_KNOWLEDGE_BASE", 200
)
KNOWLEDGE_BASE_STORAGE_QUOTA_BYTES = _positive_int_environment(
    "DEEPSEARCHER_KNOWLEDGE_BASE_STORAGE_QUOTA_BYTES", 512 * 1024 * 1024
)
TOTAL_STORAGE_QUOTA_BYTES = _positive_int_environment(
    "DEEPSEARCHER_TOTAL_STORAGE_QUOTA_BYTES", 2 * 1024 * 1024 * 1024
)
PDF_PROBE_TIMEOUT_SECONDS = _positive_int_environment("DEEPSEARCHER_PDF_PROBE_TIMEOUT_SECONDS", 10)
INGEST_REQUEST_TIMEOUT_SECONDS = _positive_int_environment(
    "DEEPSEARCHER_INGEST_REQUEST_TIMEOUT_SECONDS", 300
)
INGEST_LEASE_SECONDS = _positive_int_environment("DEEPSEARCHER_INGEST_LEASE_SECONDS", 360)
INGEST_MAX_RETRIES = _positive_int_environment("DEEPSEARCHER_INGEST_MAX_RETRIES", 3)
INGEST_RETRY_BASE_SECONDS = _positive_int_environment("DEEPSEARCHER_INGEST_RETRY_BASE_SECONDS", 5)
UPLOAD_CHUNK_BYTES = 1024 * 1024
STAGING_MAX_AGE_SECONDS = 60 * 60
EMBEDDING_BATCH_SIZE = 10
UPLOAD_DIR = DATA_DIR / "uploads"


@dataclass(frozen=True)
class StagedUpload:
    path: Path
    display_name: str
    size_bytes: int
    sha256: str


def _pipeline_config_for_document(
    session: Session,
    document: Document,
) -> tuple[list[dict], str]:
    """Reuse the latest job's pipeline config, else the default four-node chain."""
    previous = session.scalar(
        select(IngestJob)
        .where(IngestJob.document_id == document.id)
        .order_by(IngestJob.attempt.desc(), IngestJob.id.desc())
        .limit(1)
    )
    if previous is not None and previous.pipeline_steps:
        return list(previous.pipeline_steps), previous.pipeline_version or "1"
    return default_pipeline_steps(), "1"


def _connector_request_params(source_metadata: dict | None) -> dict:
    """Translate connector job metadata into load-files request overrides.

    Connector documents identify the replacement target by their stable
    Document.id (not the content sha256 used by uploads).
    """
    params: dict = {}
    replace_document_id = (source_metadata or {}).get("replace_document_id")
    if replace_document_id:
        params["replace_document_id"] = replace_document_id
    return params


def document_storage(document: Document):
    storage_type = getattr(document, "storage_type", None) or "local"
    if storage_type == "local":
        return LocalObjectStorage(UPLOAD_DIR)
    return get_object_storage(
        storage_type,
        local_root=UPLOAD_DIR,
        bucket=getattr(document, "storage_bucket", None),
    )


def document_object_key(document: Document) -> str:
    return getattr(document, "storage_key", None) or document.storage_path


def validate_document_temporal_metadata(
    *,
    published_at: date | None,
    effective_at: date | None,
    superseded_at: date | None,
) -> None:
    values = (published_at, effective_at, superseded_at)
    if any(value is not None and not isinstance(value, date) for value in values):
        raise ProductError(
            "DOCUMENT_TEMPORAL_METADATA_INVALID",
            "文档业务日期格式不正确。",
            status_code=422,
        )
    if superseded_at is not None and (
        (published_at is not None and superseded_at < published_at)
        or (effective_at is not None and superseded_at < effective_at)
    ):
        raise ProductError(
            "DOCUMENT_TEMPORAL_METADATA_INVALID",
            "文档的失效日期不能早于发布日期或生效日期。",
            status_code=422,
        )


def document_temporal_payload(document: Document) -> dict[str, str]:
    payload = {
        field: value.isoformat()
        for field in ("published_at", "effective_at", "superseded_at")
        if (value := getattr(document, field, None)) is not None
    }
    source = str(document.temporal_metadata_source or "")
    if payload and source:
        payload["temporal_metadata_source"] = source
    return payload


def validate_document_version_family(value: str | None) -> str | None:
    if value in (None, ""):
        return None
    normalized = normalize_version_family(value)
    if normalized is None:
        raise ProductError(
            "DOCUMENT_VERSION_FAMILY_INVALID",
            "文档系列标识格式不正确。",
            status_code=422,
        )
    return normalized


def document_governance_payload(document: Document) -> dict[str, str]:
    payload = document_temporal_payload(document)
    # Keep the worker tolerant of pre-0013 rows and lightweight test doubles while
    # the additive migration is rolling through an existing deployment.
    family = str(getattr(document, "version_family", None) or "")
    source = str(getattr(document, "version_family_source", None) or "")
    if family and source:
        payload["version_family"] = family
        payload["version_family_source"] = source
    return payload


def _safe_display_name(display_name: str) -> str:
    filename = Path(display_name).name.strip()
    filename = "".join(
        character for character in filename if character >= " " and character != "\x7f"
    )
    return filename[:255]


def validate_document_metadata(display_name: str, content_type: str | None) -> str:
    safe_display_name = _safe_display_name(display_name)
    if not safe_display_name:
        raise ProductError("DOCUMENT_NAME_INVALID", "文件名不能为空。")
    extension = normalize_extension(safe_display_name)
    if not extension or extension not in EXTENSION_FAMILY:
        raise ProductError(
            "DOCUMENT_UNSUPPORTED_TYPE",
            f"不支持的文件类型: .{extension or 'unknown'}",
            status_code=415,
        )
    return safe_display_name


def _make_private_directory(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)
    try:
        path.chmod(0o700)
    except OSError:
        pass


def _delete_file_quietly(path: Path) -> None:
    try:
        path.unlink(missing_ok=True)
    except OSError:
        pass


def _dispatch_ingest_or_mark_failed(
    session: Session,
    *,
    document: Document,
    job: IngestJob,
) -> None:
    try:
        dispatch_ingest_transaction(session, document=document, job=job)
    except MessageDispatchError as exc:
        now = utcnow()
        document.status = "failed"
        document.error_code = "DOCUMENT_DISPATCH_FAILED"
        document.error_message = "文档处理消息发送失败，请检查消息服务后重试。"
        job.status = "dead_letter"
        job.error_code = document.error_code
        job.error_message = document.error_message
        job.finished_at = now
        session.commit()
        raise ProductError(
            "DOCUMENT_DISPATCH_FAILED",
            "文档已保存，但处理消息发送失败，请检查消息服务后重试。",
            status_code=503,
            retryable=True,
        ) from exc


async def stage_upload(file: UploadFile) -> StagedUpload:
    display_name = validate_document_metadata(file.filename or "document.pdf", file.content_type)
    staging_dir = UPLOAD_DIR / ".staging"
    _make_private_directory(staging_dir)
    staging_path = staging_dir / f"{secrets.token_hex(24)}.part"
    digest = hashlib.sha256()
    size_bytes = 0
    prefix = bytearray()
    descriptor = os.open(staging_path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    try:
        with os.fdopen(descriptor, "wb") as destination:
            while True:
                chunk = await file.read(UPLOAD_CHUNK_BYTES)
                if not chunk:
                    break
                size_bytes += len(chunk)
                if size_bytes > MAX_PDF_BYTES:
                    raise ProductError(
                        "DOCUMENT_TOO_LARGE",
                        f"文件不能超过 {MAX_PDF_BYTES // (1024 * 1024)} MiB。",
                        status_code=413,
                    )
                if len(prefix) < 5:
                    prefix.extend(chunk[: 5 - len(prefix)])
                digest.update(chunk)
                destination.write(chunk)
            destination.flush()
            os.fsync(destination.fileno())
        if size_bytes == 0:
            raise ProductError("DOCUMENT_EMPTY", "文件不能为空。")
        try:
            validate_upload(staging_path, display_name)
        except ValueError as exc:
            raise ProductError("DOCUMENT_INVALID_CONTENT", str(exc), status_code=415) from exc
        return StagedUpload(
            path=staging_path,
            display_name=display_name,
            size_bytes=size_bytes,
            sha256=digest.hexdigest(),
        )
    except BaseException:
        _delete_file_quietly(staging_path)
        raise


async def inspect_pdf_pages(path: Path) -> int:
    process = await asyncio.create_subprocess_exec(
        sys.executable,
        "-m",
        "frontend.product.pdf_probe",
        str(path),
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.DEVNULL,
    )
    try:
        stdout, _ = await asyncio.wait_for(
            process.communicate(),
            timeout=PDF_PROBE_TIMEOUT_SECONDS,
        )
    except asyncio.TimeoutError as exc:
        process.kill()
        await process.wait()
        raise ProductError(
            "DOCUMENT_PDF_VALIDATION_TIMEOUT",
            "PDF 安全检查超时，请压缩或拆分文件后重试。",
            status_code=422,
            retryable=True,
        ) from exc
    if process.returncode != 0:
        raise ProductError(
            "DOCUMENT_INVALID_PDF",
            "PDF 结构无法解析，可能已损坏或受密码保护。",
            status_code=422,
        )
    try:
        page_count = int(stdout.decode("ascii").strip())
    except (UnicodeDecodeError, ValueError) as exc:
        raise ProductError(
            "DOCUMENT_INVALID_PDF",
            "PDF 结构无法解析，可能已损坏或受密码保护。",
            status_code=422,
        ) from exc
    if page_count < 1:
        raise ProductError("DOCUMENT_EMPTY", "PDF 文件中没有可处理的页面。", status_code=422)
    if page_count > MAX_PDF_PAGES:
        raise ProductError(
            "DOCUMENT_TOO_MANY_PAGES",
            f"PDF 页数不能超过 {MAX_PDF_PAGES} 页。",
            status_code=413,
        )
    return page_count


def _enforce_storage_quota(
    session: Session,
    *,
    knowledge_base: KnowledgeBase,
    incoming_size: int,
) -> None:
    document_count = session.scalar(
        select(func.count(Document.id)).where(Document.knowledge_base_id == knowledge_base.id)
    )
    if int(document_count or 0) >= MAX_DOCUMENTS_PER_KNOWLEDGE_BASE:
        raise ProductError(
            "KNOWLEDGE_BASE_DOCUMENT_LIMIT_REACHED",
            f"每个知识库最多保存 {MAX_DOCUMENTS_PER_KNOWLEDGE_BASE} 份文档。",
            status_code=409,
        )
    knowledge_base_usage = session.scalar(
        select(func.sum(Document.size_bytes)).where(Document.knowledge_base_id == knowledge_base.id)
    )
    if int(knowledge_base_usage or 0) + incoming_size > KNOWLEDGE_BASE_STORAGE_QUOTA_BYTES:
        raise ProductError(
            "KNOWLEDGE_BASE_STORAGE_QUOTA_EXCEEDED",
            "这个知识库的文件存储空间已满，请删除不需要的文档后重试。",
            status_code=413,
        )
    total_usage = session.scalar(select(func.sum(Document.size_bytes)))
    if int(total_usage or 0) + incoming_size > TOTAL_STORAGE_QUOTA_BYTES:
        raise ProductError(
            "DOCUMENT_STORAGE_QUOTA_EXCEEDED",
            "本地文档存储空间已满，请清理文件后重试。",
            status_code=413,
        )


def _inspect_xlsx_sheets(path: Path) -> int:
    """Visible worksheet count, using the same rule as ExcelLoader."""
    from io import BytesIO

    from openpyxl import load_workbook

    # Staging files carry a .part suffix; openpyxl validates by extension, so
    # feed the bytes through a file-like object instead.
    workbook = load_workbook(BytesIO(path.read_bytes()), read_only=True, data_only=True)
    try:
        return sum(
            1 for sheet in workbook.worksheets if sheet.sheet_state not in {"hidden", "veryHidden"}
        )
    finally:
        workbook.close()


def _inspect_pptx_slides(path: Path) -> int:
    from pptx import Presentation

    return len(Presentation(str(path)).slides)


async def inspect_document_pages(path: Path, extension: str) -> int:
    """Page-count semantics per format: pdf=pages, xlsx=visible sheets,
    pptx=slides, everything else (images and plain/office text) = 1."""
    if extension == "pdf":
        return await inspect_pdf_pages(path)
    if extension == "xlsx":
        return _inspect_xlsx_sheets(path)
    if extension == "pptx":
        return _inspect_pptx_slides(path)
    return 1


async def create_document_from_upload(
    session: Session,
    *,
    knowledge_base: KnowledgeBase,
    file: UploadFile,
    published_at: date | None = None,
    effective_at: date | None = None,
    superseded_at: date | None = None,
    version_family: str | None = None,
) -> tuple[Document, IngestJob]:
    validate_document_temporal_metadata(
        published_at=published_at,
        effective_at=effective_at,
        superseded_at=superseded_at,
    )
    normalized_version_family = validate_document_version_family(version_family)
    staged = await stage_upload(file)
    stored_key: str | None = None
    storage = get_object_storage(local_root=UPLOAD_DIR)
    try:
        page_count = await inspect_document_pages(
            staged.path,
            normalize_extension(staged.display_name),
        )
        _enforce_storage_quota(
            session,
            knowledge_base=knowledge_base,
            incoming_size=staged.size_bytes,
        )
        duplicate = session.scalar(
            select(Document).where(
                Document.knowledge_base_id == knowledge_base.id,
                Document.sha256 == staged.sha256,
            )
        )
        if duplicate is not None:
            raise ProductError(
                "DOCUMENT_DUPLICATE",
                "这个知识库中已经存在相同的 PDF。",
                status_code=409,
            )

        document = Document(
            knowledge_base_id=knowledge_base.id,
            display_name=staged.display_name,
            storage_path="",
            storage_type=storage.storage_type,
            storage_bucket=storage.bucket,
            size_bytes=staged.size_bytes,
            page_count=page_count,
            sha256=staged.sha256,
            status="queued",
            published_at=published_at,
            effective_at=effective_at,
            superseded_at=superseded_at,
            temporal_metadata_source=(
                "user_declared"
                if any(value is not None for value in (published_at, effective_at, superseded_at))
                else None
            ),
            version_family=normalized_version_family,
            version_family_source=(
                "user_declared" if normalized_version_family is not None else None
            ),
        )
        session.add(document)
        session.flush()

        # 按文件内容嗅探真实类型，落键用与之一致的扩展名：
        # 真 PDF（即使命名异常）→ pdf；markdown 内容即使被命名为 .pdf → md。
        # 避免把非 PDF 内容强制存成 .pdf，导致 core-api 用 PDFLoader 误解析。
        detected_ext = (
            detect_upload_extension(staged.path, staged.display_name)
            or normalize_extension(staged.display_name)
            or "pdf"
        )
        stored_key = storage.put_staged(
            staged.path,
            knowledge_base_id=knowledge_base.id,
            extension=detected_ext,
        )
        document.storage_key = stored_key
        # Keep the legacy field populated while callers migrate to storage_key.
        document.storage_path = stored_key
        steps, pipeline_version = _pipeline_config_for_document(session, document)
        job = create_ingest_job(
            session, document, pipeline_steps=steps, pipeline_version=pipeline_version
        )
        job.max_retries = INGEST_MAX_RETRIES
        session.commit()
        session.refresh(document)
        session.refresh(job)
        # The object and queued Document are now the durable preparation state.
        # Dispatch failure must retain them so the existing retry endpoint can recover.
        stored_key = None
        _dispatch_ingest_or_mark_failed(session, document=document, job=job)
        session.refresh(document)
        session.refresh(job)
        return document, job
    except IntegrityError as exc:
        session.rollback()
        if stored_key is not None:
            try:
                storage.delete(stored_key)
            except Exception as cleanup_exc:
                logger.warning(
                    "Failed to clean uploaded object after duplicate document rollback: "
                    "storage_type=%s bucket=%s object_key=%s error_type=%s",
                    storage.storage_type,
                    storage.bucket,
                    stored_key,
                    type(cleanup_exc).__name__,
                )
        raise ProductError(
            "DOCUMENT_DUPLICATE",
            "这个知识库中已经存在相同的 PDF。",
            status_code=409,
        ) from exc
    except Exception:
        session.rollback()
        if stored_key is not None:
            try:
                storage.delete(stored_key)
            except Exception as cleanup_exc:
                logger.warning(
                    "Failed to clean uploaded object after document rollback: "
                    "storage_type=%s bucket=%s object_key=%s error_type=%s",
                    storage.storage_type,
                    storage.bucket,
                    stored_key,
                    type(cleanup_exc).__name__,
                )
        raise
    finally:
        _delete_file_quietly(staged.path)


def cleanup_stale_uploads(*, now: float | None = None) -> int:
    staging_dir = UPLOAD_DIR / ".staging"
    if not staging_dir.is_dir():
        return 0
    cutoff = (time.time() if now is None else now) - STAGING_MAX_AGE_SECONDS
    removed = 0
    for candidate in staging_dir.glob("*.part"):
        try:
            if candidate.stat().st_mtime <= cutoff:
                candidate.unlink()
                removed += 1
        except OSError:
            continue
    return removed


def cleanup_orphaned_uploads(session: Session, *, now: float | None = None) -> int:
    if not UPLOAD_DIR.is_dir():
        return 0
    upload_root = UPLOAD_DIR.resolve()
    referenced = {
        str(Path(storage_key or storage_path).resolve())
        for storage_key, storage_path in session.execute(
            select(Document.storage_key, Document.storage_path)
        ).all()
        if storage_key or storage_path
    }
    cutoff = (time.time() if now is None else now) - STAGING_MAX_AGE_SECONDS
    removed = 0
    for candidate in UPLOAD_DIR.glob("*/*.pdf"):
        try:
            resolved = candidate.resolve()
            if upload_root not in resolved.parents or str(resolved) in referenced:
                continue
            if candidate.stat().st_mtime <= cutoff:
                candidate.unlink()
                removed += 1
        except OSError:
            continue
    return removed


class IngestProcessingError(Exception):
    def __init__(self, code: str, *, retryable: bool):
        super().__init__(code)
        self.code = code
        self.retryable = retryable


def claim_next_ingest_job(
    session: Session,
    *,
    worker_id: str,
    document_id: str | None = None,
) -> str | None:
    now = utcnow()
    query = select(IngestJob).where(
        IngestJob.status == "queued",
        or_(IngestJob.available_at.is_(None), IngestJob.available_at <= now),
    )
    if document_id is not None:
        query = query.where(IngestJob.document_id == document_id)
    query = query.order_by(IngestJob.created_at, IngestJob.id).limit(1)
    dialect_name = session.get_bind().dialect.name
    if dialect_name == "postgresql":
        query = query.with_for_update(skip_locked=True)
    candidate = session.scalar(query)
    if candidate is None:
        return None
    lease_expires_at = now + timedelta(seconds=INGEST_LEASE_SECONDS)
    if dialect_name == "postgresql":
        candidate.status = "processing"
        candidate.lease_owner = worker_id
        candidate.lease_expires_at = lease_expires_at
        candidate.retry_count += 1
        candidate.started_at = now
        candidate.finished_at = None
    else:
        result = session.execute(
            update(IngestJob)
            .where(IngestJob.id == candidate.id, IngestJob.status == "queued")
            .values(
                status="processing",
                lease_owner=worker_id,
                lease_expires_at=lease_expires_at,
                retry_count=IngestJob.retry_count + 1,
                started_at=now,
                finished_at=None,
            )
        )
        if int(result.rowcount or 0) != 1:
            session.rollback()
            return None
    document = session.get(Document, candidate.document_id)
    if document is None:
        session.rollback()
        return None
    document.status = "processing"
    document.error_code = None
    document.error_message = None
    session.commit()
    return candidate.id


async def _load_document_into_backend(
    *,
    document: Document,
    knowledge_base: KnowledgeBase,
    request_params: dict | None = None,
) -> dict:
    try:
        with document_storage(document).materialize(document_object_key(document)) as source_path:
            async with httpx.AsyncClient(
                timeout=float(INGEST_REQUEST_TIMEOUT_SECONDS),
                trust_env=False,
                headers=backend_request_headers(),
            ) as client:
                payload = {
                    "paths": str(source_path),
                    "collection_name": knowledge_base.collection_name,
                    "batch_size": EMBEDDING_BATCH_SIZE,
                    "replace_document_id": document.sha256,
                    "document_metadata": document_governance_payload(document),
                }
                if request_params:
                    payload.update(request_params)
                response = await client.post(
                    f"{BACKEND_URL}/load-files/",
                    json=payload,
                )
    except FileNotFoundError as exc:
        raise IngestProcessingError("DOCUMENT_CONTENT_MISSING", retryable=False) from exc
    except StorageError as exc:
        raise IngestProcessingError("DOCUMENT_STORAGE_INVALID", retryable=False) from exc
    except (httpx.ConnectTimeout, httpx.PoolTimeout) as exc:
        raise IngestProcessingError("DOCUMENT_PROCESSING_UNAVAILABLE", retryable=True) from exc
    except (httpx.ReadTimeout, httpx.WriteTimeout) as exc:
        # The server may still be mutating the index after a response timeout. Do
        # not start a concurrent retry; move the task to manual reconciliation.
        raise IngestProcessingError("DOCUMENT_PROCESSING_TIMEOUT", retryable=False) from exc
    except httpx.ConnectError as exc:
        raise IngestProcessingError("DOCUMENT_PROCESSING_UNAVAILABLE", retryable=True) from exc
    except httpx.RequestError as exc:
        raise IngestProcessingError(
            "DOCUMENT_PROCESSING_TRANSPORT_FAILED", retryable=False
        ) from exc
    if not response.is_success:
        retryable = response.status_code in {429, 502, 503, 504} or response.status_code >= 500
        raise IngestProcessingError("DOCUMENT_PROCESSING_FAILED", retryable=retryable)
    try:
        response_payload = response.json()
    except ValueError as exc:
        raise IngestProcessingError("DOCUMENT_PROCESSING_INVALID_RESPONSE", retryable=True) from exc
    manifest = (
        response_payload.get("collection", {}).get("manifest")
        if isinstance(response_payload, dict)
        else None
    )
    if not isinstance(manifest, dict):
        raise IngestProcessingError("DOCUMENT_INDEX_MANIFEST_MISSING", retryable=True)
    return manifest


def _finish_ingest_job(
    *,
    job_id: str,
    worker_id: str,
    manifest: dict | None = None,
    failure: IngestProcessingError | None = None,
    failure_detail: dict | None = None,
) -> None:
    with SessionLocal() as session:
        job = session.get(IngestJob, job_id)
        if job is None or job.status != "processing" or job.lease_owner != worker_id:
            return
        document = session.get(Document, job.document_id)
        if document is None:
            return
        knowledge_base = session.get(KnowledgeBase, document.knowledge_base_id)
        if knowledge_base is None:
            return
        now = utcnow()
        detail_text = None
        if failure_detail:
            try:
                detail_text = json.dumps(failure_detail, ensure_ascii=False)[:280]
            except (TypeError, ValueError):
                detail_text = None
        job.lease_owner = None
        job.lease_expires_at = None
        if failure is None and manifest is not None:
            knowledge_base.index_manifest = json.dumps(
                manifest,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            )
            document.status = "ready"
            document.error_code = None
            document.error_message = None
            job.status = "succeeded"
            job.error_code = None
            job.error_message = None
            job.finished_at = now
        elif failure is not None and failure.retryable and job.retry_count < job.max_retries:
            delay = min(
                INGEST_RETRY_BASE_SECONDS * (2 ** max(job.retry_count - 1, 0)),
                60,
            )
            document.status = "queued"
            document.error_code = "DOCUMENT_RETRY_SCHEDULED"
            document.error_message = "文档处理暂时失败，系统将在后台自动重试。"
            job.status = "queued"
            job.available_at = now + timedelta(seconds=delay)
            job.error_code = failure.code
            job.error_message = detail_text or "文档处理暂时失败，任务已进入重试队列。"
            job.finished_at = None
        else:
            document.status = "failed"
            document.error_code = (
                failure.code if failure is not None else "DOCUMENT_PROCESSING_FAILED"
            )
            document.error_message = detail_text or "文档处理失败，请确认文件和问答服务状态后重试。"
            job.status = "dead_letter"
            job.error_code = document.error_code
            job.error_message = document.error_message
            job.finished_at = now
        knowledge_base.updated_at = now
        session.commit()


def _finalize_ingest(
    context: IngestionContext,
    *,
    job_id: str,
    worker_id: str,
) -> None:
    """Single owner of _finish_ingest_job after a pipeline run (P2-B).

    Nodes never mutate job/document final state; lifecycle success requires
    context.manifest to have been produced.
    """
    if context.manifest is not None:
        _finish_ingest_job(job_id=job_id, worker_id=worker_id, manifest=context.manifest)
        return
    if context.failure is not None:
        _finish_ingest_job(
            job_id=job_id,
            worker_id=worker_id,
            failure=IngestProcessingError(
                context.failure.code,
                retryable=context.failure.retryable,
            ),
            failure_detail=context.failure.to_dict(),
        )
        return
    # Terminated (or completed) without a manifest is NOT a successful ingestion.
    _finish_ingest_job(
        job_id=job_id,
        worker_id=worker_id,
        failure=IngestProcessingError("DOCUMENT_PROCESSING_INCOMPLETE", retryable=False),
        failure_detail={
            "code": "DOCUMENT_PROCESSING_INCOMPLETE",
            "message": "流水线未产生入库结果",
        },
    )


async def process_claimed_ingest_job(job_id: str, *, worker_id: str) -> None:
    with SessionLocal() as session:
        job = session.get(IngestJob, job_id)
        if job is None or job.status != "processing" or job.lease_owner != worker_id:
            return
        document = session.get(Document, job.document_id)
        if document is None:
            return
        knowledge_base = session.get(KnowledgeBase, document.knowledge_base_id)
        if knowledge_base is None:
            return
        steps = list(job.pipeline_steps) if job.pipeline_steps else None
        source_metadata = dict(job.source_metadata or {})
        expected_hash = source_metadata.get("content_hash")
        if expected_hash and document.content_hash != expected_hash:
            # A stale enqueued job (the source was updated/deleted since) must not
            # re-index old bytes; drop it without touching the current document.
            job.status = "dead_letter"
            job.error_code = "DOCUMENT_STALE"
            job.error_message = "源文档版本已更新，旧入库任务已丢弃。"
            job.finished_at = utcnow()
            session.commit()
            return
        session.expunge(document)
        session.expunge(knowledge_base)
    context = IngestionContext(
        job=job,
        document=document,
        knowledge_base=knowledge_base,
        steps=steps or default_pipeline_steps(),
        request_params=_connector_request_params(source_metadata),
    )
    try:
        await execute_chain(context)
    except Exception as exc:
        context.status = STATUS_FAILED
        context.failure = context.failure or NodeFailure(
            node_id="chain",
            node_type="pipeline",
            code="PIPELINE_UNEXPECTED",
            message=str(exc),
            retryable=True,
        )
    _finalize_ingest(context, job_id=job_id, worker_id=worker_id)


async def process_rocketmq_ingest_job(
    job_id: str,
    *,
    worker_id: str,
    delivery_attempt: int,
    lease_seconds: int = INGEST_LEASE_SECONDS,
) -> bool:
    """Process one MQ delivery and return whether the consumer should ACK it."""

    with SessionLocal() as session:
        job = session.get(IngestJob, job_id)
        if job is None or job.status in {"succeeded", "dead_letter"}:
            return True
        now = utcnow()
        if job.status != "processing":
            return False
        lease_expires_at = job.lease_expires_at
        if lease_expires_at is not None and lease_expires_at.tzinfo is None:
            lease_expires_at = lease_expires_at.replace(tzinfo=now.tzinfo)
        if (
            job.lease_owner is not None
            and job.lease_owner != worker_id
            and lease_expires_at is not None
            and lease_expires_at > now
        ):
            return False
        document = session.get(Document, job.document_id)
        if document is None:
            return True
        knowledge_base = session.get(KnowledgeBase, document.knowledge_base_id)
        if knowledge_base is None:
            return True
        job.lease_owner = worker_id
        job.lease_expires_at = now + timedelta(seconds=lease_seconds)
        job.retry_count = max(job.retry_count + 1, delivery_attempt)
        job.started_at = now
        max_retries = job.max_retries
        session.commit()
        steps = list(job.pipeline_steps) if job.pipeline_steps else None
        source_metadata = dict(job.source_metadata or {})
        expected_hash = source_metadata.get("content_hash")
        if expected_hash and document.content_hash != expected_hash:
            job.status = "dead_letter"
            job.error_code = "DOCUMENT_STALE"
            job.error_message = "源文档版本已更新，旧入库任务已丢弃。"
            job.finished_at = utcnow()
            session.commit()
            return True
        session.expunge(document)
        session.expunge(knowledge_base)
    context = IngestionContext(
        job=job,
        document=document,
        knowledge_base=knowledge_base,
        steps=steps or default_pipeline_steps(),
        request_params=_connector_request_params(source_metadata),
    )
    try:
        await execute_chain(context)
    except Exception as exc:
        context.status = STATUS_FAILED
        context.failure = context.failure or NodeFailure(
            node_id="chain",
            node_type="pipeline",
            code="PIPELINE_UNEXPECTED",
            message=str(exc),
            retryable=True,
        )

    if context.manifest is not None:
        _finish_ingest_job(job_id=job_id, worker_id=worker_id, manifest=context.manifest)
        return True

    failure = IngestProcessingError(
        context.failure.code if context.failure is not None else "DOCUMENT_PROCESSING_INCOMPLETE",
        retryable=context.failure.retryable if context.failure is not None else False,
    )
    if failure.retryable and delivery_attempt < max_retries:
        with SessionLocal() as session:
            current_job = session.get(IngestJob, job_id)
            current_document = (
                session.get(Document, current_job.document_id) if current_job is not None else None
            )
            if current_job is not None and current_job.lease_owner == worker_id:
                current_job.lease_owner = None
                current_job.lease_expires_at = None
                current_job.error_code = failure.code
                current_job.error_message = "文档处理暂时失败，等待消息重试。"
                if current_document is not None:
                    current_document.status = "processing"
                    current_document.error_code = "DOCUMENT_RETRY_SCHEDULED"
                    current_document.error_message = "文档处理暂时失败，系统将在后台自动重试。"
                session.commit()
        return False

    _finish_ingest_job(
        job_id=job_id,
        worker_id=worker_id,
        failure=failure,
        failure_detail=context.failure.to_dict() if context.failure is not None else None,
    )
    return True


async def process_document(document_id: str) -> None:
    """Compatibility helper for tests and direct callers; the product uses the durable worker."""
    worker_id = f"inline-{os.getpid()}-{uuid4().hex[:12]}"
    with SessionLocal() as session:
        job_id = claim_next_ingest_job(
            session,
            worker_id=worker_id,
            document_id=document_id,
        )
    if job_id is not None:
        await process_claimed_ingest_job(job_id, worker_id=worker_id)


def retry_document(session: Session, document: Document) -> IngestJob:
    if document.status != "failed":
        raise ProductError(
            "DOCUMENT_NOT_RETRYABLE",
            "只有处理失败的文档可以重试。",
            status_code=409,
        )
    document.status = "queued"
    document.error_code = None
    document.error_message = None
    steps, pipeline_version = _pipeline_config_for_document(session, document)
    job = create_ingest_job(
        session,
        document,
        pipeline_steps=steps,
        pipeline_version=pipeline_version,
    )
    job.max_retries = INGEST_MAX_RETRIES
    session.commit()
    session.refresh(job)
    _dispatch_ingest_or_mark_failed(session, document=document, job=job)
    session.refresh(document)
    session.refresh(job)
    return job


def update_document_temporal_metadata(
    session: Session,
    document: Document,
    *,
    published_at: date | None,
    effective_at: date | None,
    superseded_at: date | None,
) -> IngestJob | None:
    validate_document_temporal_metadata(
        published_at=published_at,
        effective_at=effective_at,
        superseded_at=superseded_at,
    )
    if document.status == "processing":
        raise ProductError(
            "DOCUMENT_BUSY",
            "文档正在处理，完成后才能修改业务日期。",
            status_code=409,
            retryable=True,
        )
    values = (published_at, effective_at, superseded_at)
    if values == (document.published_at, document.effective_at, document.superseded_at):
        return None
    document.published_at = published_at
    document.effective_at = effective_at
    document.superseded_at = superseded_at
    document.temporal_metadata_source = (
        "user_declared" if any(value is not None for value in values) else None
    )
    job = None
    if document.status == "ready":
        document.status = "queued"
        document.error_code = None
        document.error_message = None
        steps, pipeline_version = _pipeline_config_for_document(session, document)
        job = create_ingest_job(
            session, document, pipeline_steps=steps, pipeline_version=pipeline_version
        )
        job.max_retries = INGEST_MAX_RETRIES
    document.knowledge_base.updated_at = utcnow()
    session.commit()
    session.refresh(document)
    if job is not None:
        session.refresh(job)
        _dispatch_ingest_or_mark_failed(session, document=document, job=job)
        session.refresh(document)
        session.refresh(job)
    return job


def update_document_governance_metadata(
    session: Session,
    document: Document,
    *,
    published_at: date | None,
    effective_at: date | None,
    superseded_at: date | None,
    version_family: str | None,
) -> IngestJob | None:
    validate_document_temporal_metadata(
        published_at=published_at,
        effective_at=effective_at,
        superseded_at=superseded_at,
    )
    normalized_family = validate_document_version_family(version_family)
    if document.status == "processing":
        raise ProductError(
            "DOCUMENT_BUSY",
            "文档正在处理，完成后才能修改治理元数据。",
            status_code=409,
            retryable=True,
        )
    values = (published_at, effective_at, superseded_at, normalized_family)
    current = (
        document.published_at,
        document.effective_at,
        document.superseded_at,
        document.version_family,
    )
    if values == current:
        return None
    document.published_at = published_at
    document.effective_at = effective_at
    document.superseded_at = superseded_at
    document.temporal_metadata_source = (
        "user_declared"
        if any(value is not None for value in (published_at, effective_at, superseded_at))
        else None
    )
    document.version_family = normalized_family
    document.version_family_source = "user_declared" if normalized_family is not None else None
    job = None
    if document.status == "ready":
        document.status = "queued"
        document.error_code = None
        document.error_message = None
        steps, pipeline_version = _pipeline_config_for_document(session, document)
        job = create_ingest_job(
            session, document, pipeline_steps=steps, pipeline_version=pipeline_version
        )
        job.max_retries = INGEST_MAX_RETRIES
    document.knowledge_base.updated_at = utcnow()
    session.commit()
    session.refresh(document)
    if job is not None:
        session.refresh(job)
        _dispatch_ingest_or_mark_failed(session, document=document, job=job)
        session.refresh(document)
        session.refresh(job)
    return job


async def delete_document(session: Session, document: Document) -> None:
    if document.status in {"queued", "processing"}:
        raise ProductError(
            "DOCUMENT_BUSY",
            "文档正在处理，完成后才能删除。",
            status_code=409,
            retryable=True,
        )

    knowledge_base = session.get(KnowledgeBase, document.knowledge_base_id)
    if knowledge_base is None:
        raise ProductError(
            "KNOWLEDGE_BASE_NOT_FOUND",
            "没有找到这个知识库。",
            status_code=404,
        )

    try:
        async with httpx.AsyncClient(
            timeout=30.0,
            trust_env=False,
            headers=backend_request_headers(),
        ) as client:
            response = await client.delete(
                f"{BACKEND_URL}/collections/{knowledge_base.collection_name}"
                f"/documents/{document.sha256}"
            )
        if not response.is_success:
            raise ProductError(
                "DOCUMENT_VECTOR_DELETE_FAILED",
                "向量数据删除失败，文档尚未删除，请稍后重试。",
                status_code=502,
                retryable=True,
            )
        try:
            response_payload = response.json()
        except (AttributeError, ValueError):
            response_payload = {}
        manifest = response_payload.get("manifest") if isinstance(response_payload, dict) else None
        if isinstance(manifest, dict):
            knowledge_base.index_manifest = json.dumps(
                manifest,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            )
    except httpx.RequestError as exc:
        raise ProductError(
            "DOCUMENT_VECTOR_SERVICE_UNAVAILABLE",
            "问答服务暂时不可用，文档尚未删除，请稍后重试。",
            status_code=503,
            retryable=True,
        ) from exc

    try:
        document_storage(document).delete(document_object_key(document))
    except (OSError, StorageError) as exc:
        raise ProductError(
            "DOCUMENT_FILE_DELETE_FAILED",
            "文档对象删除失败，请检查存储服务后重试。",
            status_code=500,
            retryable=True,
        ) from exc

    knowledge_base.updated_at = utcnow()
    session.delete(document)
    session.commit()
