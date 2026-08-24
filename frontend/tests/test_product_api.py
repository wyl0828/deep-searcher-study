import asyncio
import base64
import hashlib
import io
import json
from contextlib import contextmanager
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import sessionmaker

from deepsearcher.provenance import build_trust_provenance
from deepsearcher.trust import build_temporal_context
from frontend.product.auth import require_user
from frontend.product.db import Base, create_database_engine, get_session
from frontend.product.errors import ProductError
from frontend.product.models import (
    AnswerRun,
    Citation,
    Conversation,
    Department,
    DepartmentKnowledgeBase,
    Document,
    KnowledgeBase,
    KnowledgeHealthSnapshot,
    Message,
    MessageFeedback,
    User,
)
from frontend.product.repositories import get_conversation
from frontend.product.services import conversations as conversation_service
from frontend.product.services import documents as document_service
from frontend.product.services.conversations import (
    _finish_assistant_message,
    _iter_sse_events,
    _validate_stream_envelope,
    build_conversation_history,
    collect_supported_citations,
    query_error_from_response,
    stream_message_events,
)
from frontend.product.services.query_scope import resolve_conversation_scope
from frontend.server import app


@pytest.fixture(autouse=True)
def stub_pdf_page_inspection(monkeypatch):
    async def inspect_pdf_pages(_path):
        return 1

    monkeypatch.setattr(
        "frontend.product.services.documents.inspect_pdf_pages",
        inspect_pdf_pages,
    )


@contextmanager
def product_client(tmp_path):
    engine = create_database_engine(f"sqlite:///{(tmp_path / 'api.db').as_posix()}")
    Base.metadata.create_all(engine)
    test_session = sessionmaker(bind=engine, expire_on_commit=False)
    with test_session() as session:
        test_user = User(
            username="test-user",
            display_name="测试用户",
            password_hash="disabled",
            role="admin",
        )
        session.add(test_user)
        session.commit()

    def override_session():
        with test_session() as session:
            yield session

    app.dependency_overrides[get_session] = override_session
    app.dependency_overrides[require_user] = lambda: test_user
    try:
        client = TestClient(app)
        client.product_session_factory = test_session
        client.admin_user = test_user
        yield client
    finally:
        app.dependency_overrides.clear()


def parse_sse_events(payload: str) -> list[dict]:
    events = []
    for frame in payload.replace("\r\n", "\n").split("\n\n"):
        data_lines = [line[5:].lstrip() for line in frame.splitlines() if line.startswith("data:")]
        if data_lines:
            events.append(json.loads("\n".join(data_lines)))
    return events


def _route_payload(mode: str = "knowledge") -> dict:
    intent = {
        "chat": "general_question",
        "knowledge": "enterprise_fact",
        "web": "external_fact",
    }[mode]
    return {
        "answer_mode": mode,
        "route_intent": intent,
        "route_source": "rule",
        "reason_code": "test_route",
        "confidence": 0.95,
        "router_version": "query-router-v1",
        "initial_risk_level": "medium",
        "initial_risk_factors": [],
    }


class _RouteResponse:
    is_success = True

    def __init__(self, mode: str = "knowledge"):
        self._payload = _route_payload(mode)

    def json(self):
        return self._payload


def test_stream_envelope_requires_matching_request_id_and_contiguous_sequence():
    first = {
        "version": 1,
        "request_id": "request.with-dots-1",
        "sequence": 1,
        "event": "started",
        "data": {"stage": "query_started"},
    }

    assert _validate_stream_envelope(
        "started",
        first,
        expected_request_id="request.with-dots-1",
        last_sequence=0,
    ) == ("request.with-dots-1", 1)

    for invalid in (
        {**first, "request_id": "another-request"},
        {**first, "sequence": 2},
        {**first, "version": 2},
        {**first, "event": "routing"},
    ):
        with pytest.raises(ProductError) as exc_info:
            _validate_stream_envelope(
                "started",
                invalid,
                expected_request_id="request.with-dots-1",
                last_sequence=0,
            )
        assert exc_info.value.code == "QUERY_STREAM_INVALID"


def test_stream_parser_rejects_oversized_frames(monkeypatch):
    monkeypatch.setattr(conversation_service, "MAX_SSE_LINE_CHARS", 32)

    class Response:
        async def aiter_lines(self):
            yield "event: started"
            yield "data: " + ("x" * 33)
            yield ""

    async def consume():
        return [event async for event in _iter_sse_events(Response())]

    with pytest.raises(ProductError) as exc_info:
        asyncio.run(consume())
    assert exc_info.value.code == "QUERY_STREAM_INVALID"


def test_citations_never_bind_to_a_document_from_another_knowledge_base(tmp_path):
    with product_client(tmp_path) as client:
        own_kb = client.post(
            "/api/knowledge-bases",
            json={"name": "当前知识库", "description": ""},
        ).json()
        foreign_kb = client.post(
            "/api/knowledge-bases",
            json={"name": "其他知识库", "description": ""},
        ).json()
        conversation_payload = client.post(
            "/api/conversations",
            json={"knowledge_base_id": own_kb["id"]},
        ).json()

        with client.product_session_factory() as session:
            foreign_document = Document(
                knowledge_base_id=foreign_kb["id"],
                display_name="foreign-secret.pdf",
                storage_path=str(tmp_path / "foreign-secret.pdf"),
                size_bytes=10,
                sha256="a" * 64,
                status="ready",
            )
            assistant_message = Message(
                conversation_id=conversation_payload["id"],
                role="assistant",
                content="answer",
                status="pending",
            )
            session.add_all([foreign_document, assistant_message])
            session.commit()
            trace = {
                "iterations": [
                    {
                        "retrieved_documents": [
                            {
                                "document_id": foreign_document.id,
                                "display_name": "safe-fallback.pdf",
                                "text": "bounded evidence",
                                "supported": True,
                            }
                        ]
                    }
                ]
            }

            citations = collect_supported_citations(
                session,
                message=assistant_message,
                trace=trace,
            )

            assert len(citations) == 1
            assert citations[0].document_id is None
            assert citations[0].display_name == "safe-fallback.pdf"
            assert "foreign-secret" not in citations[0].display_name


def test_claim_grounding_persists_partial_support_and_rejects_fake_evidence_ids(tmp_path):
    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "声明引用库", "description": ""},
        ).json()
        conversation = client.post(
            "/api/conversations",
            json={"knowledge_base_id": knowledge_base["id"]},
        ).json()

        with client.product_session_factory() as session:
            document = Document(
                knowledge_base_id=knowledge_base["id"],
                display_name="grounding.pdf",
                storage_path=str(tmp_path / "grounding.pdf"),
                size_bytes=10,
                sha256="e" * 64,
                status="ready",
            )
            assistant = Message(
                conversation_id=conversation["id"],
                role="assistant",
                content="",
                status="pending",
            )
            session.add_all([document, assistant])
            session.commit()
            _finish_assistant_message(
                session,
                assistant_message=assistant,
                payload={
                    "result": "Milvus 是向量数据库。[E1] 它支持任意 SQL。[E9]",
                    "trace": {
                        "grounding": {
                            "version": 1,
                            "state": "partially_grounded",
                            "evidence": [
                                {
                                    "evidence_id": "E1",
                                    "document_id": document.id,
                                    "display_name": "grounding.pdf",
                                    "page_number": 2,
                                    "location_id": "loc-grounding-1",
                                    "text": "Milvus 是向量数据库。",
                                    "supported": True,
                                }
                            ],
                            "claims": [
                                {
                                    "index": 1,
                                    "text": "Milvus 是向量数据库。",
                                    "status": "supported",
                                    "evidence_ids": ["E1"],
                                    "citation_spans": [
                                        {
                                            "evidence_id": "E1",
                                            "start": 0,
                                            "end": len("Milvus 是向量数据库。"),
                                            "text": "Milvus 是向量数据库。",
                                            "match_type": "normalized_exact",
                                            "score": 1.0,
                                        },
                                        {
                                            "evidence_id": "E1",
                                            "start": 0,
                                            "end": 6,
                                            "text": "伪造片段",
                                            "match_type": "normalized_exact",
                                            "score": 1.0,
                                        },
                                    ],
                                },
                                {
                                    "index": 2,
                                    "text": "它支持任意 SQL。",
                                    "status": "invalid_citation",
                                    "evidence_ids": [],
                                    "invalid_evidence_ids": ["E9"],
                                    "citation_spans": [
                                        {
                                            "evidence_id": "E9",
                                            "start": 0,
                                            "end": 999999,
                                            "text": "不得持久化",
                                            "match_type": "normalized_exact",
                                            "score": 1.0,
                                        }
                                    ],
                                },
                            ],
                        }
                    },
                },
            )

        detail = client.get(f"/api/conversations/{conversation['id']}").json()
        persisted = detail["messages"][0]
        assert persisted["answer_state"] == "partially_grounded"
        assert len(persisted["citations"]) == 1
        assert persisted["claims"] == [
            {
                "id": persisted["claims"][0]["id"],
                "index": 1,
                "text": "Milvus 是向量数据库。",
                "support_status": "supported",
                "structural_support_status": "supported",
                "citation_indices": [1],
                "citation_spans": [
                    {
                        "citation_index": 1,
                        "start": 0,
                        "end": len("Milvus 是向量数据库。"),
                        "text": "Milvus 是向量数据库。",
                        "match_type": "normalized_exact",
                        "score": 1.0,
                    }
                ],
                "citation_status": "valid",
                "entailment_status": "not_checked",
                "consistency_status": "not_checked",
                "consistency_checks": [],
                "risk_status": "not_assessed",
                "risk_checks": [],
                "confidence": None,
                "reason_codes": ["CITATION_VALID"],
            },
            {
                "id": persisted["claims"][1]["id"],
                "index": 2,
                "text": "它支持任意 SQL。",
                "support_status": "invalid_citation",
                "structural_support_status": "invalid_citation",
                "citation_indices": [],
                "citation_spans": [],
                "citation_status": "invalid",
                "entailment_status": "not_checked",
                "consistency_status": "not_checked",
                "consistency_checks": [],
                "risk_status": "not_assessed",
                "risk_checks": [],
                "confidence": None,
                "reason_codes": ["CITATION_INVALID"],
            },
        ]


def test_structured_grounding_without_claims_never_falls_back_to_grounded(tmp_path):
    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "空声明引用库", "description": ""},
        ).json()
        conversation = client.post(
            "/api/conversations",
            json={"knowledge_base_id": knowledge_base["id"]},
        ).json()

        with client.product_session_factory() as session:
            assistant = Message(
                conversation_id=conversation["id"],
                role="assistant",
                content="",
                status="pending",
            )
            session.add(assistant)
            session.commit()
            _finish_assistant_message(
                session,
                assistant_message=assistant,
                payload={
                    "result": "```python\nprint(42)\n```",
                    "trace": {
                        "grounding": {
                            "version": 1,
                            "state": "insufficient_evidence",
                            "evidence": [
                                {
                                    "evidence_id": "E1",
                                    "display_name": "code.pdf",
                                    "page_number": 1,
                                    "text": "print(42)",
                                    "supported": True,
                                }
                            ],
                            "claims": [],
                        }
                    },
                },
            )

        persisted = client.get(f"/api/conversations/{conversation['id']}").json()["messages"][0]
        assert persisted["answer_state"] == "insufficient_evidence"
        assert len(persisted["citations"]) == 1
        assert persisted["claims"] == []


def test_trust_contract_and_policy_decision_are_persisted(tmp_path):
    provenance = build_trust_provenance(
        SimpleNamespace(
            config=SimpleNamespace(
                provide_settings={
                    "llm": {"provider": "TestLLM", "config": {"model": "test-model"}},
                    "embedding": {
                        "provider": "TestEmbedding",
                        "config": {"model": "test-embedding"},
                    },
                },
                query_settings={},
                load_settings={},
            ),
            llm=SimpleNamespace(model="test-model"),
            embedding_model=None,
            vector_db=None,
            entailment_checker=None,
        ),
        evidence_snapshot=[],
        execution_scope="online",
        temporal_context=build_temporal_context(
            reference_time="2026-08-11T03:00:00+00:00",
            timezone_name="Asia/Shanghai",
        ),
    )
    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "可信策略库", "description": ""},
        ).json()
        conversation = client.post(
            "/api/conversations",
            json={"knowledge_base_id": knowledge_base["id"]},
        ).json()

        with client.product_session_factory() as session:
            assistant = Message(
                conversation_id=conversation["id"],
                role="assistant",
                content="",
                status="pending",
            )
            session.add(assistant)
            session.commit()
            _finish_assistant_message(
                session,
                assistant_message=assistant,
                payload={
                    "result": "根据现有证据，无法给出有充分依据的回答。",
                    "trace": {
                        "grounding": {
                            "version": 1,
                            "state": "insufficient_evidence",
                            "evidence": [],
                            "claims": [],
                        },
                        "trust": {
                            "version": 1,
                            "trust_status": "insufficient_evidence",
                            "safety_status": "not_evaluated",
                            "verification_level": "semantic_entailment_partial",
                            "evidence_snapshot_available": True,
                            "entailment": {
                                "version": 1,
                                "checker": "llm_nli",
                                "checker_version": "1.0.0",
                                "status": "partial",
                                "token_usage": 17,
                                "eligible_claim_count": 1,
                                "exact_match_count": 0,
                                "checker_claim_count": 1,
                                "entailed_count": 0,
                                "contradicted_count": 0,
                                "unknown_count": 1,
                                "not_checked_count": 0,
                                "error_code": None,
                            },
                            "risk": {
                                "version": 1,
                                "classifier": "deterministic_query_risk",
                                "classifier_version": "1.1.0",
                                "risk_level": "high",
                                "query_type": "financial_policy",
                                "risk_factors": [
                                    "FINANCIAL_DOMAIN",
                                    "DECISION_REQUEST",
                                    "QUANTITATIVE_DECISION",
                                ],
                                "requirements": {
                                    "require_citation": True,
                                    "require_decisive_entailment": True,
                                    "minimum_evidence_count": 2,
                                    "minimum_distinct_source_count": 2,
                                    "allow_unknown_entailment": False,
                                },
                            },
                            "freshness": {
                                "version": 1,
                                "classifier": "deterministic_freshness_intent",
                                "classifier_version": "1.1.0",
                                "required": True,
                                "mode": "latest_effective",
                                "ordering_basis": "effective_at",
                                "reason_codes": ["FRESHNESS_LATEST_EFFECTIVE_REQUESTED"],
                            },
                            "temporal_context": {
                                "version": 1,
                                "source": "request_clock",
                                "reference_time": "2026-08-11T03:00:00+00:00",
                                "reference_date": "2026-08-11",
                                "timezone": "Asia/Shanghai",
                            },
                            "input": {
                                "trust_status": "insufficient_evidence",
                                "claim_count": 1,
                                "supported_claim_count": 0,
                                "claims": [
                                    {
                                        "index": 1,
                                        "text": "每份 PDF 最大 100 MB。",
                                        "support_status": "unsupported",
                                        "structural_support_status": "supported",
                                        "citation_status": "valid",
                                        "entailment_status": "not_checked",
                                        "entailment_method": "semantic_nli",
                                        "entailment_checker": "llm_nli",
                                        "entailment_checker_version": "1.0.0",
                                        "confidence": 0.61,
                                        "risk_status": "rejected",
                                        "risk_checks": [
                                            {
                                                "kind": "entailment",
                                                "status": "failed",
                                                "reason_code": "HIGH_RISK_ENTAILMENT_REQUIRED",
                                                "actual": "unknown",
                                                "required": "entailed_or_contradicted",
                                            }
                                        ],
                                        "consistency_status": "inconsistent",
                                        "consistency_checks": [
                                            {
                                                "kind": "quantity",
                                                "status": "inconsistent",
                                                "reason_code": "QUANTITY_NOT_IN_EVIDENCE",
                                                "claim_values": ["100|mb"],
                                                "missing_values": ["100|mb"],
                                            }
                                        ],
                                        "reason_codes": [
                                            "CITATION_VALID",
                                            "QUANTITY_NOT_IN_EVIDENCE",
                                        ],
                                    }
                                ],
                            },
                            "output": {
                                "trust_status": "insufficient_evidence",
                                "claim_count": 1,
                                "supported_claim_count": 0,
                                "claims": [],
                            },
                            "claims": [],
                            "policy": {
                                "profile": "strict_high_risk",
                                "risk_level": "high",
                                "action": "refuse",
                                "reason_codes": ["EVIDENCE_INSUFFICIENT"],
                            },
                            "limitations": [
                                "SEMANTIC_ENTAILMENT_NOT_CHECKED",
                                "SAFETY_NOT_EVALUATED",
                            ],
                            "provenance": provenance,
                        },
                    },
                },
            )

        persisted = client.get(f"/api/conversations/{conversation['id']}").json()["messages"][0]
        assert persisted["trust_contract_version"] == 1
        assert persisted["trust_status"] == "insufficient_evidence"
        assert persisted["safety_status"] == "not_evaluated"
        assert persisted["policy_action"] == "refuse"
        assert persisted["policy_profile"] == "strict_high_risk"
        assert persisted["policy_reason_codes"] == ["EVIDENCE_INSUFFICIENT"]
        assert persisted["risk_level"] == "high"
        assert persisted["query_type"] == "financial_policy"
        assert persisted["risk_factors"] == [
            "FINANCIAL_DOMAIN",
            "DECISION_REQUEST",
            "QUANTITATIVE_DECISION",
        ]
        assert persisted["provenance_contract_version"] == 2
        assert persisted["provenance_digest"] == provenance["digest"]
        assert persisted["trust_details"]["provenance"] == provenance
        assert persisted["trust_details"]["provenance"]["evidence"] == {
            "snapshot_status": "complete",
            "evidence_count": 0,
            "knowledge_base_count": 0,
            "web_count": 0,
            "snapshot_fingerprint": provenance["evidence"]["snapshot_fingerprint"],
            "items": [],
        }
        assert persisted["trust_details"]["verification_level"] == ("semantic_entailment_partial")
        assert persisted["trust_details"]["entailment"] == {
            "version": 1,
            "checker": "llm_nli",
            "checker_version": "1.0.0",
            "status": "partial",
            "token_usage": 17,
            "eligible_claim_count": 1,
            "exact_match_count": 0,
            "checker_claim_count": 1,
            "entailed_count": 0,
            "contradicted_count": 0,
            "unknown_count": 1,
            "not_checked_count": 0,
        }
        assert persisted["trust_details"]["risk"]["requirements"] == {
            "require_citation": True,
            "require_decisive_entailment": True,
            "minimum_evidence_count": 2,
            "minimum_distinct_source_count": 2,
            "allow_unknown_entailment": False,
        }
        assert persisted["trust_details"]["freshness"] == {
            "version": 1,
            "classifier": "deterministic_freshness_intent",
            "classifier_version": "1.1.0",
            "required": True,
            "mode": "latest_effective",
            "ordering_basis": "effective_at",
            "reason_codes": ["FRESHNESS_LATEST_EFFECTIVE_REQUESTED"],
        }
        assert persisted["trust_details"]["temporal_context"] == {
            "version": 1,
            "source": "request_clock",
            "reference_time": "2026-08-11T03:00:00+00:00",
            "reference_date": "2026-08-11",
            "timezone": "Asia/Shanghai",
        }
        finding = persisted["trust_details"]["input"]["claims"][0]
        assert finding["structural_support_status"] == "supported"
        assert finding["support_status"] == "unsupported"
        assert finding["consistency_checks"][0]["missing_values"] == ["100|mb"]
        assert finding["entailment_method"] == "semantic_nli"
        assert finding["entailment_checker"] == "llm_nli"
        assert finding["confidence"] == 0.61
        assert finding["risk_status"] == "rejected"
        assert finding["risk_checks"][0]["actual"] == "unknown"


def test_create_list_and_select_knowledge_bases(tmp_path):
    with product_client(tmp_path) as client:
        first = client.post(
            "/api/knowledge-bases",
            json={"name": "AI 学习资料", "description": "课程与项目资料"},
        )
        second = client.post(
            "/api/knowledge-bases",
            json={"name": "产品文档", "description": ""},
        )

        assert first.status_code == 201
        assert first.json()["is_current"] is True
        assert second.status_code == 201
        assert second.json()["is_current"] is False

        selected = client.put(f"/api/knowledge-bases/{second.json()['id']}/current")
        assert selected.status_code == 200
        assert selected.json()["is_current"] is True

        items = client.get("/api/knowledge-bases").json()["items"]
        assert items[0]["name"] == "产品文档"
        assert sum(item["is_current"] for item in items) == 1


def test_product_errors_use_safe_structured_shape(tmp_path):
    with product_client(tmp_path) as client:
        response = client.get("/api/knowledge-bases/kb_missing")

    assert response.status_code == 404
    payload = response.json()
    assert payload["error"].pop("request_id")
    assert payload == {
        "error": {
            "code": "KNOWLEDGE_BASE_NOT_FOUND",
            "message": "没有找到这个知识库。",
            "retryable": False,
        }
    }


def test_document_upload_rejects_unknown_extension(tmp_path):
    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "资料库", "description": ""},
        ).json()
        response = client.post(
            f"/api/knowledge-bases/{knowledge_base['id']}/documents",
            files={"file": ("note.xyz", b"whatever", "application/octet-stream")},
        )

    assert response.status_code == 415
    assert response.json()["error"]["code"] == "DOCUMENT_UNSUPPORTED_TYPE"


def test_document_upload_rejects_container_mismatch(tmp_path):
    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "资料库", "description": ""},
        ).json()
        response = client.post(
            f"/api/knowledge-bases/{knowledge_base['id']}/documents",
            files={"file": ("notes.md", b"\x00\x01\x02\x03binary", "text/plain")},
        )

    assert response.status_code == 415
    assert response.json()["error"]["code"] == "DOCUMENT_INVALID_CONTENT"


def test_document_upload_accepts_xlsx_pptx_and_image(tmp_path):
    from openpyxl import Workbook
    from pptx import Presentation

    workbook = Workbook()
    for sheet_name in ("收入", "费用"):
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(["项目", "金额"])
        sheet.append(["A", 100])
    workbook.remove(workbook["Sheet"])
    xlsx_buffer = io.BytesIO()
    workbook.save(xlsx_buffer)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "季度汇报"
    pptx_buffer = io.BytesIO()
    presentation.save(pptx_buffer)

    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "多格式库", "description": ""},
        ).json()
        kb_id = knowledge_base["id"]

        xlsx = client.post(
            f"/api/knowledge-bases/{kb_id}/documents",
            files={
                "file": (
                    "table.xlsx",
                    xlsx_buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert xlsx.status_code == 202
        assert xlsx.json()["document"]["page_count"] == 2  # visible sheets

        pptx = client.post(
            f"/api/knowledge-bases/{kb_id}/documents",
            files={
                "file": (
                    "deck.pptx",
                    pptx_buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )
            },
        )
        assert pptx.status_code == 202
        assert pptx.json()["document"]["page_count"] == 1  # slides

        png = client.post(
            f"/api/knowledge-bases/{kb_id}/documents",
            files={"file": ("image.png", PNG_BYTES, "image/png")},
        )
        assert png.status_code == 202
        assert png.json()["document"]["page_count"] == 1


def test_document_upload_persists_explicit_business_dates(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "frontend.product.services.documents.UPLOAD_DIR",
        tmp_path / "uploads",
    )
    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "制度资料", "description": ""},
        ).json()
        response = client.post(
            f"/api/knowledge-bases/{knowledge_base['id']}/documents",
            files={"file": ("policy.pdf", b"%PDF-1.7\npolicy", "application/pdf")},
            data={
                "published_at": "2026-08-11",
                "effective_at": "2026-08-12",
                "superseded_at": "2027-01-01",
                "version_family": "Travel Expense Policy",
            },
        )

    assert response.status_code == 202
    document = response.json()["document"]
    assert document["published_at"] == "2026-08-11"
    assert document["effective_at"] == "2026-08-12"
    assert document["superseded_at"] == "2027-01-01"
    assert document["temporal_metadata_source"] == "user_declared"
    assert document["version_family"] == "travel-expense-policy"
    assert document["version_family_source"] == "user_declared"


def test_document_upload_rejects_superseded_date_before_effective_date(tmp_path):
    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "制度资料", "description": ""},
        ).json()
        response = client.post(
            f"/api/knowledge-bases/{knowledge_base['id']}/documents",
            files={"file": ("policy.pdf", b"%PDF-1.7\npolicy", "application/pdf")},
            data={
                "effective_at": "2026-08-12",
                "superseded_at": "2026-08-11",
            },
        )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "DOCUMENT_TEMPORAL_METADATA_INVALID"


def test_updating_ready_document_business_dates_queues_metadata_reindex(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(
        "frontend.product.services.documents.UPLOAD_DIR",
        tmp_path / "uploads",
    )
    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "制度资料", "description": ""},
        ).json()
        upload = client.post(
            f"/api/knowledge-bases/{knowledge_base['id']}/documents",
            files={"file": ("policy.pdf", b"%PDF-1.7\npolicy", "application/pdf")},
        ).json()
        document_id = upload["document"]["id"]
        with client.product_session_factory() as session:
            stored = session.get(Document, document_id)
            assert stored is not None
            stored.status = "ready"
            session.commit()

        response = client.patch(
            f"/api/documents/{document_id}/governance-metadata",
            json={
                "published_at": "2026-08-11",
                "effective_at": "2026-08-12",
                "superseded_at": None,
                "version_family": "Travel Expense Policy",
            },
        )

    assert response.status_code == 202
    payload = response.json()
    assert payload["document"]["status"] == "queued"
    assert payload["document"]["published_at"] == "2026-08-11"
    assert payload["document"]["version_family"] == "travel-expense-policy"
    assert payload["document"]["version_family_source"] == "user_declared"
    assert payload["job"]["attempt"] == 2


def test_document_worker_sends_temporal_metadata_to_core_ingest(tmp_path, monkeypatch):
    upload_root = tmp_path / "uploads"
    source_dir = upload_root / "kb"
    source_dir.mkdir(parents=True)
    source = source_dir / "policy.pdf"
    source.write_bytes(b"%PDF-1.7\npolicy")
    monkeypatch.setattr(document_service, "UPLOAD_DIR", upload_root)
    captured = {}

    class Response:
        is_success = True

        @staticmethod
        def json():
            return {"collection": {"manifest": {"schema_version": 2}}}

    class Client:
        def __init__(self, **_kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def post(self, url, *, json):
            captured["url"] = url
            captured["json"] = json
            return Response()

    monkeypatch.setattr(document_service.httpx, "AsyncClient", Client)
    document = SimpleNamespace(
        storage_path=str(source),
        sha256="a" * 64,
        published_at=date(2026, 8, 11),
        effective_at=date(2026, 8, 12),
        superseded_at=None,
        temporal_metadata_source="user_declared",
        version_family="travel-expense-policy",
        version_family_source="user_declared",
    )
    knowledge_base = SimpleNamespace(collection_name="kb_" + "a" * 32)

    manifest = asyncio.run(
        document_service._load_document_into_backend(
            document=document,
            knowledge_base=knowledge_base,
        )
    )

    assert manifest == {"schema_version": 2}
    assert captured["json"]["document_metadata"] == {
        "published_at": "2026-08-11",
        "effective_at": "2026-08-12",
        "temporal_metadata_source": "user_declared",
        "version_family": "travel-expense-policy",
        "version_family_source": "user_declared",
    }


def test_document_content_returns_inline_original_pdf(tmp_path, monkeypatch):
    upload_dir = tmp_path / "uploads"
    monkeypatch.setattr("frontend.product.services.documents.UPLOAD_DIR", upload_dir)
    pdf_payload = b"%PDF-1.7\ninline source"

    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "原文资料", "description": ""},
        ).json()
        upload = client.post(
            f"/api/knowledge-bases/{knowledge_base['id']}/documents",
            files={"file": ("guide.pdf", pdf_payload, "application/pdf")},
        ).json()

        response = client.get(f"/api/documents/{upload['document']['id']}/content")

    assert response.status_code == 200
    assert response.content == pdf_payload
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].startswith('inline; filename="guide.pdf"')
    assert response.headers["cache-control"] == "private, max-age=3600"
    assert response.headers["x-content-type-options"] == "nosniff"


def test_document_content_rejects_storage_path_outside_upload_root(tmp_path, monkeypatch):
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    monkeypatch.setattr("frontend.product.services.documents.UPLOAD_DIR", upload_dir)
    outside_path = tmp_path / "outside.pdf"
    outside_path.write_bytes(b"%PDF-1.7\noutside")

    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "异常路径资料", "description": ""},
        ).json()
        with client.product_session_factory() as session:
            document = Document(
                knowledge_base_id=knowledge_base["id"],
                display_name="outside.pdf",
                storage_path=str(outside_path),
                size_bytes=outside_path.stat().st_size,
                sha256="f" * 64,
                status="ready",
            )
            session.add(document)
            session.commit()
            document_id = document.id

        response = client.get(f"/api/documents/{document_id}/content")

    assert response.status_code == 500
    assert response.json()["error"]["code"] == "DOCUMENT_STORAGE_INVALID"


def test_delete_document_removes_vectors_file_and_database_record(tmp_path, monkeypatch):
    upload_dir = tmp_path / "uploads"
    monkeypatch.setattr("frontend.product.services.documents.UPLOAD_DIR", upload_dir)
    captured = {}
    deleted_manifest = {
        "schema_version": 1,
        "embedding_model": "embedding-a",
        "data_version": "after-delete",
    }

    class Response:
        is_success = True

        @staticmethod
        def json():
            return {"delete_count": 2, "manifest": deleted_manifest}

    class Client:
        def __init__(self, **kwargs):
            captured["client_kwargs"] = kwargs

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def delete(self, url, *, headers=None):
            captured["url"] = url
            captured["headers"] = headers
            return Response()

    monkeypatch.setattr("frontend.product.services.documents.httpx.AsyncClient", Client)

    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "待清理资料", "description": ""},
        ).json()
        upload = client.post(
            f"/api/knowledge-bases/{knowledge_base['id']}/documents",
            files={"file": ("guide.pdf", b"%PDF-1.7\ndelete me", "application/pdf")},
        ).json()
        document_id = upload["document"]["id"]

        with client.product_session_factory() as session:
            document = session.get(Document, document_id)
            assert document is not None
            document.status = "ready"
            storage_path = document.storage_path
            source_hash = document.sha256
            conversation = Conversation(knowledge_base_id=knowledge_base["id"])
            message = Message(
                conversation=conversation,
                role="assistant",
                content="带历史引用的回答",
                status="succeeded",
            )
            citation = Citation(
                message=message,
                document_id=document.id,
                index=1,
                display_name=document.display_name,
                page_number=1,
                text="历史引用文字",
            )
            session.add(conversation)
            session.commit()
            citation_id = citation.id

        response = client.delete(f"/api/documents/{document_id}")

        assert response.status_code == 204
        assert client.get(f"/api/documents/{document_id}").status_code == 404
        detail = client.get(f"/api/knowledge-bases/{knowledge_base['id']}").json()
        assert detail["index_manifest"] == deleted_manifest
        with client.product_session_factory() as session:
            retained_citation = session.get(Citation, citation_id)
            assert retained_citation is not None
            assert retained_citation.document_id is None
            assert retained_citation.display_name == "guide.pdf"
            assert retained_citation.text == "历史引用文字"

    assert not Path(storage_path).exists()
    assert source_hash in captured["url"]
    assert captured["url"].endswith(source_hash)
    assert captured["client_kwargs"] == {
        "timeout": 30.0,
        "trust_env": False,
        "headers": {"X-DeepSearcher-Tenant": "local"},
    }
    assert storage_path


def test_delete_document_blocks_processing_document(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "frontend.product.services.documents.UPLOAD_DIR",
        tmp_path / "uploads",
    )
    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "处理中资料", "description": ""},
        ).json()
        upload = client.post(
            f"/api/knowledge-bases/{knowledge_base['id']}/documents",
            files={"file": ("guide.pdf", b"%PDF-1.7\nbusy", "application/pdf")},
        ).json()

        response = client.delete(f"/api/documents/{upload['document']['id']}")

        assert response.status_code == 409
        assert response.json()["error"]["code"] == "DOCUMENT_BUSY"
        assert client.get(f"/api/documents/{upload['document']['id']}").status_code == 200


def test_delete_document_keeps_local_record_when_vector_delete_fails(tmp_path, monkeypatch):
    upload_dir = tmp_path / "uploads"
    monkeypatch.setattr("frontend.product.services.documents.UPLOAD_DIR", upload_dir)

    class Response:
        is_success = False

    class Client:
        def __init__(self, **_kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def delete(self, _url):
            return Response()

    monkeypatch.setattr("frontend.product.services.documents.httpx.AsyncClient", Client)

    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "删除失败资料", "description": ""},
        ).json()
        upload = client.post(
            f"/api/knowledge-bases/{knowledge_base['id']}/documents",
            files={"file": ("guide.pdf", b"%PDF-1.7\nkeep me", "application/pdf")},
        ).json()
        document_id = upload["document"]["id"]

        with client.product_session_factory() as session:
            document = session.get(Document, document_id)
            assert document is not None
            document.status = "ready"
            storage_path = document.storage_path
            session.commit()

        response = client.delete(f"/api/documents/{document_id}")

        assert response.status_code == 502
        assert response.json()["error"]["code"] == "DOCUMENT_VECTOR_DELETE_FAILED"
        assert client.get(f"/api/documents/{document_id}").status_code == 200
        assert Path(storage_path).exists()
        assert storage_path


def test_delete_knowledge_base_cascades_data_files_vectors_and_selects_fallback(
    tmp_path,
    monkeypatch,
):
    upload_dir = tmp_path / "uploads"
    monkeypatch.setattr("frontend.product.services.documents.UPLOAD_DIR", upload_dir)
    captured = {}

    class Response:
        is_success = True

    class Client:
        def __init__(self, **kwargs):
            captured["client_kwargs"] = kwargs

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def delete(self, url, *, headers=None):
            captured["url"] = url
            captured["headers"] = headers
            return Response()

    monkeypatch.setattr(
        "frontend.product.services.knowledge_bases.httpx.AsyncClient",
        Client,
    )

    with product_client(tmp_path) as client:
        deleting = client.post(
            "/api/knowledge-bases",
            json={"name": "待删除知识库", "description": ""},
        ).json()
        fallback = client.post(
            "/api/knowledge-bases",
            json={"name": "回退知识库", "description": ""},
        ).json()
        upload = client.post(
            f"/api/knowledge-bases/{deleting['id']}/documents",
            files={"file": ("guide.pdf", b"%PDF-1.7\ncascade", "application/pdf")},
        ).json()
        document_id = upload["document"]["id"]

        with client.product_session_factory() as session:
            stored_knowledge_base = session.get(KnowledgeBase, deleting["id"])
            assert stored_knowledge_base is not None
            collection_name = stored_knowledge_base.collection_name
            document = session.get(Document, document_id)
            assert document is not None
            document.status = "ready"
            conversation = Conversation(
                knowledge_base_id=deleting["id"],
                title="待级联对话",
            )
            message = Message(
                conversation=conversation,
                role="assistant",
                content="待级联回答",
                status="succeeded",
            )
            citation = Citation(
                message=message,
                document_id=document.id,
                index=1,
                display_name=document.display_name,
                text="待级联引用",
            )
            session.add(conversation)
            session.commit()
            conversation_id = conversation.id
            message_id = message.id
            citation_id = citation.id

        detail = client.get(f"/api/knowledge-bases/{deleting['id']}").json()
        assert detail["document_count"] == 1
        assert detail["conversation_count"] == 1

        response = client.delete(f"/api/knowledge-bases/{deleting['id']}")

        assert response.status_code == 200
        assert response.json() == {
            "deleted_id": deleting["id"],
            "current_knowledge_base_id": fallback["id"],
        }
        assert client.get(f"/api/knowledge-bases/{deleting['id']}").status_code == 404
        assert client.get(f"/api/documents/{document_id}").status_code == 404
        assert client.get(f"/api/conversations/{conversation_id}").status_code == 404
        assert client.get(f"/api/knowledge-bases/{fallback['id']}").json()["is_current"]
        with client.product_session_factory() as session:
            assert session.get(Message, message_id) is None
            assert session.get(Citation, citation_id) is None

    assert not upload_dir.joinpath(deleting["id"]).exists()
    assert collection_name in captured["url"]
    assert captured["url"].endswith(collection_name)
    assert captured["headers"] == {"X-Confirm-Collection": collection_name}
    assert captured["client_kwargs"] == {
        "timeout": 30.0,
        "trust_env": False,
        "headers": {"X-DeepSearcher-Tenant": "local"},
    }


def test_delete_knowledge_base_blocks_while_document_is_processing(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(
        "frontend.product.services.documents.UPLOAD_DIR",
        tmp_path / "uploads",
    )
    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "处理中知识库", "description": ""},
        ).json()
        client.post(
            f"/api/knowledge-bases/{knowledge_base['id']}/documents",
            files={"file": ("guide.pdf", b"%PDF-1.7\nbusy", "application/pdf")},
        )

        response = client.delete(f"/api/knowledge-bases/{knowledge_base['id']}")

        assert response.status_code == 409
        assert response.json()["error"]["code"] == "KNOWLEDGE_BASE_BUSY"
        assert client.get(f"/api/knowledge-bases/{knowledge_base['id']}").status_code == 200


def test_reindex_knowledge_base_sends_all_sources_and_persists_manifest(
    tmp_path,
    monkeypatch,
):
    upload_dir = tmp_path / "uploads"
    monkeypatch.setattr("frontend.product.services.documents.UPLOAD_DIR", upload_dir)
    captured = {}
    manifest = {
        "schema_version": 1,
        "embedding_provider": "TestEmbedding",
        "embedding_model": "embedding-a",
        "embedding_version": "embedding-a-v1",
        "dimension": 8,
        "normalization": "none",
        "metric_type": "L2",
        "chunk_config_version": "chunk-v1",
        "document_version": "document-v1",
        "data_version": "data-v1",
    }

    class Response:
        is_success = True

        @staticmethod
        def json():
            return {
                "collection": {
                    "manifest": manifest,
                    "previous_collection": "kb_previous_version",
                }
            }

    class Client:
        def __init__(self, **kwargs):
            captured["client_kwargs"] = kwargs

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def post(self, url, *, headers, json):
            captured["url"] = url
            captured["headers"] = headers
            captured["json"] = json
            return Response()

    monkeypatch.setattr(
        "frontend.product.services.knowledge_bases.httpx.AsyncClient",
        Client,
    )

    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "待重建知识库", "description": "完整资料"},
        ).json()
        upload = client.post(
            f"/api/knowledge-bases/{knowledge_base['id']}/documents",
            files={"file": ("guide.pdf", b"%PDF-1.7\nreindex", "application/pdf")},
        ).json()
        with client.product_session_factory() as session:
            document = session.get(Document, upload["document"]["id"])
            assert document is not None
            document.status = "failed"
            document.error_code = "DOCUMENT_PROCESSING_FAILED"
            document.error_message = "旧索引处理失败"
            source_path = document.storage_path
            collection_name = document.knowledge_base.collection_name
            session.commit()

        response = client.post(f"/api/knowledge-bases/{knowledge_base['id']}/reindex")

        assert response.status_code == 200
        assert response.json()["index_status"] == "verified"
        assert response.json()["index_manifest"] == manifest
        assert response.json()["index_previous_collection"] == "kb_previous_version"
        assert response.json()["ready_document_count"] == 1
        listed = client.get("/api/knowledge-bases").json()["items"][0]
        assert listed["index_status"] == "verified"
        assert listed["index_manifest"] == manifest
        with client.product_session_factory() as session:
            rebuilt_document = session.get(Document, upload["document"]["id"])
            assert rebuilt_document is not None
            assert rebuilt_document.status == "ready"
            assert rebuilt_document.error_code is None
            assert rebuilt_document.error_message is None

    assert captured["url"].endswith(f"/collections/{collection_name}/rebuild")
    assert captured["headers"] == {"X-Confirm-Collection": collection_name}
    assert captured["json"] == {
        "paths": [source_path],
        "document_metadata": [{}],
        "collection_description": "完整资料",
        "batch_size": 10,
    }
    assert captured["client_kwargs"] == {
        "timeout": 600.0,
        "trust_env": False,
        "headers": {"X-DeepSearcher-Tenant": "local"},
    }


def test_delete_conversation_cascades_messages_and_citations_only(tmp_path):
    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "保留知识库", "description": ""},
        ).json()
        conversation = client.post(
            "/api/conversations",
            json={"knowledge_base_id": knowledge_base["id"]},
        ).json()

        with client.product_session_factory() as session:
            stored_conversation = session.get(Conversation, conversation["id"])
            message = Message(
                conversation=stored_conversation,
                role="assistant",
                content="待删除回答",
                status="succeeded",
            )
            citation = Citation(
                message=message,
                index=1,
                display_name="历史资料.pdf",
                text="待删除引用",
            )
            session.add(message)
            session.commit()
            message_id = message.id
            citation_id = citation.id

        response = client.delete(f"/api/conversations/{conversation['id']}")

        assert response.status_code == 204
        assert client.get(f"/api/conversations/{conversation['id']}").status_code == 404
        assert client.get(f"/api/knowledge-bases/{knowledge_base['id']}").status_code == 200
        with client.product_session_factory() as session:
            assert session.get(Message, message_id) is None
            assert session.get(Citation, citation_id) is None


def test_conversation_history_excludes_failed_and_weakly_grounded_answers():
    conversation = Conversation(knowledge_base_id="kb_test", title="上下文测试")
    conversation.messages = [
        Message(role="user", content="Milvus 有哪些部署方式？", status="succeeded"),
        Message(
            role="assistant",
            content="包括单机部署和集群部署。",
            status="succeeded",
            answer_state="fully_grounded",
        ),
        Message(role="user", content="它们有什么区别？", status="succeeded"),
        Message(
            role="assistant",
            content="这是一条只有部分依据的回答。",
            status="succeeded",
            answer_state="partially_grounded",
        ),
        Message(role="user", content="再说说扩缩容。", status="succeeded"),
        Message(
            role="assistant",
            content="调用失败的错误文字。",
            status="failed",
            answer_state="failed",
        ),
    ]

    assert build_conversation_history(conversation) == [
        {
            "role": "user",
            "content": "Milvus 有哪些部署方式？",
            "grounded": False,
        },
        {
            "role": "assistant",
            "content": "包括单机部署和集群部署。",
            "grounded": True,
        },
        {"role": "user", "content": "它们有什么区别？", "grounded": False},
        {"role": "user", "content": "再说说扩缩容。", "grounded": False},
    ]


def test_chat_completion_does_not_create_grounding_state_or_citations(tmp_path):
    """Public Chat prose must remain outside the RAG evidence contract."""

    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "Chat 回归库", "description": ""},
        ).json()
        conversation = client.post(
            "/api/conversations",
            json={"knowledge_base_id": knowledge_base["id"]},
        ).json()

        with client.product_session_factory() as session:
            assistant = Message(
                conversation_id=conversation["id"],
                role="assistant",
                content="",
                status="pending",
                answer_mode="chat",
            )
            session.add(assistant)
            session.commit()
            _finish_assistant_message(
                session,
                assistant_message=assistant,
                payload={
                    "result": "这是通用回答。",
                    "trace": {
                        "grounding": {
                            "version": 1,
                            "evidence": [{"text": "不应成为引用", "supported": True}],
                        }
                    },
                },
            )
            assert assistant.answer_state is None
            assert assistant.trust_status == "not_assessed"
            assert assistant.citations == []
            assert assistant.claims == []


def test_upload_query_and_citation_flow_uses_internal_collection_scope(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(
        "frontend.product.services.documents.UPLOAD_DIR",
        tmp_path / "uploads",
    )
    pdf_payload = b"%PDF-1.7\nproduct test"
    source_hash = hashlib.sha256(pdf_payload).hexdigest()
    captured = {}

    class Response:
        is_success = True

        @staticmethod
        def json():
            return {
                "result": "DeepSearcher 会先理解问题，再检索并生成答案。",
                "consume_token": 42,
                "trace": {
                    "version": 1,
                    "iterations": [
                        {
                            "retrieved_documents": [
                                {
                                    "document_id": source_hash,
                                    "display_name": "guide.pdf",
                                    "page_number": 2,
                                    "chunk_index": 3,
                                    "section_title": "查询流程",
                                    "section_path": ["使用指南", "查询流程"],
                                    "char_start": 120,
                                    "char_end": 148,
                                    "bbox": [0.1, 0.2, 0.9, 0.3],
                                    "location_id": "location-query-flow",
                                    "source_locator": "page=2&char=120-148",
                                    "parser_version": "pdfplumber-layout-v2+rapidocr-v3",
                                    "extraction_method": "text",
                                    "source_type": "knowledge_base",
                                    "source_url": None,
                                    "source_domain": None,
                                    "trusted": True,
                                    "text": "DeepSearcher 查询流程说明。",
                                    "supported": True,
                                }
                            ]
                        }
                    ],
                },
            }

    class Client:
        def __init__(self, **kwargs):
            captured["client_kwargs"] = kwargs

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def post(self, url, json):
            if url.endswith("/route"):
                captured["route_url"] = url
                captured["route_json"] = json
                return _RouteResponse("knowledge")
            captured["url"] = url
            captured["json"] = json
            return Response()

    monkeypatch.setattr(
        "frontend.product.services.conversations.httpx.AsyncClient",
        Client,
    )

    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "AI 全栈学习资料", "description": ""},
        ).json()
        upload = client.post(
            f"/api/knowledge-bases/{knowledge_base['id']}/documents",
            files={"file": ("guide.pdf", pdf_payload, "application/pdf")},
        )

        assert upload.status_code == 202
        product_document_id = upload.json()["document"]["id"]
        assert upload.json()["document"]["status"] == "queued"
        job_id = upload.json()["job"]["id"]
        persisted_job = client.get(f"/api/ingest-jobs/{job_id}")
        assert persisted_job.status_code == 200
        assert persisted_job.json()["document_id"] == product_document_id
        assert persisted_job.json()["status"] == "queued"

        # Frozen scope contract: a fixed conversation can query only after the
        # document has become ready and the index manifest is verified.
        with client.product_session_factory() as session:
            session.get(Document, product_document_id).status = "ready"
            session.get(KnowledgeBase, knowledge_base["id"]).index_manifest = json.dumps(
                {"schema_version": 1}
            )
            session.commit()

        conversation = client.post(
            "/api/conversations",
            json={"knowledge_base_id": knowledge_base["id"]},
        ).json()
        answer = client.post(
            f"/api/conversations/{conversation['id']}/messages",
            headers={"X-Request-ID": "product-query-test-1"},
            json={"content": "DeepSearcher 的查询流程是什么？"},
        )

        assert answer.status_code == 200
        assistant = answer.json()["assistant_message"]
        assert assistant["status"] == "succeeded"
        assert assistant["answer_state"] == "grounded"
        assert assistant["citations"] == [
            {
                "id": assistant["citations"][0]["id"],
                "index": 1,
                "document_id": product_document_id,
                "display_name": "guide.pdf",
                "page_number": 2,
                "chunk_index": 3,
                "section_title": "查询流程",
                "section_path": ["使用指南", "查询流程"],
                "char_start": 120,
                "char_end": 148,
                "bbox": [0.1, 0.2, 0.9, 0.3],
                "location_id": "location-query-flow",
                "source_locator": "page=2&char=120-148",
                "parser_version": "pdfplumber-layout-v2+rapidocr-v3",
                "extraction_method": "text",
                "source_type": "knowledge_base",
                "source_url": None,
                "source_domain": None,
                "trusted": True,
                "published_at": None,
                "effective_at": None,
                "superseded_at": None,
                "temporal_metadata_source": None,
                "version_family": None,
                "version_family_source": None,
                "text": "DeepSearcher 查询流程说明。",
                "supported": True,
            }
        ]

        detail = client.get(f"/api/conversations/{conversation['id']}").json()
        assert len(detail["messages"]) == 2
        persisted_citation = detail["messages"][1]["citations"][0]
        assert persisted_citation["page_number"] == 2
        assert persisted_citation["section_title"] == "查询流程"
        assert persisted_citation["location_id"] == "location-query-flow"

    assert captured["url"].endswith("/query")
    assert captured["route_url"].endswith("/route")
    assert captured["route_json"]["use_web_search"] is False
    assert captured["json"]["collection_names"][0].startswith("kb_")
    assert captured["json"]["original_query"] == "DeepSearcher 的查询流程是什么？"
    assert captured["json"]["conversation_history"] == []
    assert captured["json"]["include_trace"] is True
    assert captured["client_kwargs"]["trust_env"] is False
    assert captured["client_kwargs"]["headers"]["X-Request-ID"] == "product-query-test-1"


def test_product_query_stream_relays_safe_stages_and_persists_completion(
    tmp_path,
    monkeypatch,
):
    captured = {}
    upstream_events = [
        {
            "version": 1,
            "request_id": "product-stream-test-1",
            "sequence": 1,
            "event": "started",
            "data": {
                "stage": "query_started",
                "api_key": "should-not-reach-browser",
                "path": r"C:\Users\private\notes.pdf",
            },
        },
        {
            "version": 1,
            "request_id": "product-stream-test-1",
            "sequence": 2,
            "event": "contextualization",
            "data": {
                "depends_on_history": True,
                "history_turn_count": 4,
                "fallback_used": False,
                "reason": "rewritten",
                "token_usage": 19,
                "standalone_query": "private rewritten query",
            },
        },
        {
            "version": 1,
            "request_id": "product-stream-test-1",
            "sequence": 3,
            "event": "routing",
            "data": {
                "agent": "DeepSearch",
                "fallback_used": False,
                "subquery": "private intermediate query",
            },
        },
        {
            "version": 1,
            "request_id": "product-stream-test-1",
            "sequence": 4,
            "event": "retrieval",
            "data": {
                "iteration": 1,
                "retrieved_count": 4,
                "metadata": {"source_path": r"C:\Users\private\notes.pdf"},
            },
        },
        {
            "version": 1,
            "request_id": "product-stream-test-1",
            "sequence": 5,
            "event": "web_search",
            "data": {
                "iteration": 1,
                "status": "completed",
                "provider": "tavily",
                "query_count": 1,
                "result_count": 1,
                "error_code": None,
                "api_key": "should-not-reach-browser",
            },
        },
        {
            "version": 1,
            "request_id": "product-stream-test-1",
            "sequence": 6,
            "event": "support",
            "data": {"iteration": 1, "supported_count": 1},
        },
        {
            "version": 1,
            "request_id": "product-stream-test-1",
            "sequence": 7,
            "event": "reflection",
            "data": {"iteration": 1, "has_enough_information": True},
        },
        {
            "version": 1,
            "request_id": "product-stream-test-1",
            "sequence": 8,
            "event": "completed",
            "data": {
                "result": "这是基于资料生成的安全回答。",
                "trace": {
                    "version": 3,
                    "iterations": [
                        {
                            "retrieved_documents": [
                                {
                                    "display_name": "guide.pdf",
                                    "page_number": 2,
                                    "text": "可核对的证据片段。",
                                    "supported": True,
                                },
                                {
                                    "display_name": "Tavily 官方文档",
                                    "reference": "https://docs.tavily.com/search",
                                    "source_type": "web",
                                    "source_url": "https://docs.tavily.com/search",
                                    "source_domain": "docs.tavily.com",
                                    "trusted": True,
                                    "text": "可核对的网页证据片段。",
                                    "supported": True,
                                },
                            ]
                        }
                    ],
                    "private_debug": "must-not-reach-browser",
                },
            },
        },
    ]

    class Response:
        is_success = True

        async def aiter_lines(self):
            for envelope in upstream_events:
                yield f"event: {envelope['event']}"
                yield f"data: {json.dumps(envelope)}"
                yield ""

    class StreamContext:
        async def __aenter__(self):
            return Response()

        async def __aexit__(self, *_args):
            return None

    class Client:
        def __init__(self, **kwargs):
            captured["client_kwargs"] = kwargs

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def post(self, url, json):
            assert url.endswith("/route")
            captured["route_url"] = url
            captured["route_json"] = json
            return _RouteResponse("web")

        def stream(self, method, url, *, json):
            captured["method"] = method
            captured["url"] = url
            captured["json"] = json
            return StreamContext()

    monkeypatch.setattr(
        "frontend.product.services.conversations.httpx.AsyncClient",
        Client,
    )

    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "实时回答知识库", "description": ""},
        ).json()
        with client.product_session_factory() as session:
            session.add(
                Document(
                    knowledge_base_id=knowledge_base["id"],
                    display_name="stream-guide.pdf",
                    storage_path=str(tmp_path / "stream-guide.pdf"),
                    size_bytes=10,
                    sha256="s" * 64,
                    status="ready",
                )
            )
            session.get(KnowledgeBase, knowledge_base["id"]).index_manifest = json.dumps(
                {"schema_version": 1}
            )
            session.commit()
        conversation = client.post(
            "/api/conversations",
            json={"knowledge_base_id": knowledge_base["id"]},
        ).json()

        response = client.post(
            f"/api/conversations/{conversation['id']}/messages/stream",
            headers={"X-Request-ID": "product-stream-test-1"},
            json={"content": "请给出可核对的回答", "use_web_search": True},
        )

        assert response.status_code == 200
        assert response.headers["cache-control"] == "no-store, no-transform"
        assert response.headers["x-accel-buffering"] == "no"
        assert response.headers["x-trace-retention"] == "transient"
        assert response.headers["x-request-id"] == "product-stream-test-1"
        events = parse_sse_events(response.text)
        assert [event["event"] for event in events] == [
            "started",
            "contextualization",
            "routing",
            "retrieval",
            "web_search",
            "support",
            "reflection",
            "completed",
        ]
        assert events[1]["data"] == {
            "depends_on_history": True,
            "history_turn_count": 4,
            "fallback_used": False,
            "reason": "rewritten",
        }
        assert events[3]["data"] == {"iteration": 1, "retrieved_count": 4}
        assert events[4]["data"] == {
            "iteration": 1,
            "status": "completed",
            "provider": "tavily",
            "query_count": 1,
            "result_count": 1,
            "error_code": None,
        }
        assert set(events[-1]["data"]) == {"user_message", "assistant_message"}
        assert events[-1]["data"]["assistant_message"]["content"] == (
            "这是基于资料生成的安全回答。"
        )
        assert events[-1]["data"]["assistant_message"]["citations"][0]["text"] == (
            "可核对的证据片段。"
        )
        for private_value in (
            "should-not-reach-browser",
            r"C:\Users\private\notes.pdf",
            "private intermediate query",
            "private rewritten query",
            "must-not-reach-browser",
            "private_debug",
        ):
            assert private_value not in response.text

        detail = client.get(f"/api/conversations/{conversation['id']}").json()
        assert [message["status"] for message in detail["messages"]] == [
            "succeeded",
            "succeeded",
        ]
        assert detail["messages"][1]["citations"][0]["display_name"] == "guide.pdf"
        web_citation = detail["messages"][1]["citations"][1]
        assert web_citation["source_type"] == "web"
        assert web_citation["source_url"] == "https://docs.tavily.com/search"
        assert web_citation["source_domain"] == "docs.tavily.com"
        assert web_citation["trusted"] is True
        assert web_citation["document_id"] is None

    assert captured["method"] == "POST"
    assert captured["route_url"].endswith("/route")
    assert captured["route_json"]["use_web_search"] is True
    assert captured["url"].endswith("/query/stream")
    assert captured["json"]["collection_names"][0].startswith("kb_")
    assert captured["json"]["retrieval_mode"] == "web"
    assert "use_web_search" not in captured["json"]
    assert captured["client_kwargs"]["trust_env"] is False
    assert captured["client_kwargs"]["headers"]["X-Request-ID"] == "product-stream-test-1"


def test_closing_product_query_stream_marks_pending_answer_as_stopped(
    tmp_path,
    monkeypatch,
):
    upstream_closed = {"value": False}
    started = {
        "version": 1,
        "request_id": "request-cancel-1",
        "sequence": 1,
        "event": "started",
        "data": {"stage": "query_started"},
    }

    class Response:
        is_success = True

        async def aiter_lines(self):
            yield "event: started"
            yield f"data: {json.dumps(started)}"
            yield ""

    class StreamContext:
        async def __aenter__(self):
            return Response()

        async def __aexit__(self, *_args):
            upstream_closed["value"] = True

    class Client:
        def __init__(self, **_kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def post(self, url, json):
            assert url.endswith("/route")
            return _RouteResponse("knowledge")

        def stream(self, *_args, **_kwargs):
            return StreamContext()

    monkeypatch.setattr(
        "frontend.product.services.conversations.httpx.AsyncClient",
        Client,
    )

    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "中断回答知识库", "description": ""},
        ).json()
        conversation_payload = client.post(
            "/api/conversations",
            json={"knowledge_base_id": knowledge_base["id"]},
        ).json()

        with client.product_session_factory() as session:
            conversation = get_conversation(session, conversation_payload["id"])
            assert conversation is not None

            async def consume_then_close():
                stream = stream_message_events(
                    session,
                    conversation=conversation,
                    content="开始后立即停止",
                )
                first = await anext(stream)
                await stream.aclose()
                return first

            first_event = asyncio.run(consume_then_close())

        assert first_event["event"] == "started"
        assert upstream_closed["value"] is True
        detail = client.get(f"/api/conversations/{conversation_payload['id']}").json()
        assert [message["status"] for message in detail["messages"]] == [
            "succeeded",
            "failed",
        ]
        assert detail["messages"][1]["content"] == "本次回答已停止。"


def test_query_error_mapping_preserves_vector_failure_categories():
    class Response:
        def __init__(self, code):
            self.code = code

        def json(self):
            return {"error": {"code": self.code}}

    cases = [
        ("RUNTIME_INITIALIZATION_FAILED", 503, True, "正在恢复依赖"),
        ("VECTOR_DB_UNAVAILABLE", 503, True, "暂时不可用"),
        ("VECTOR_COLLECTION_NOT_FOUND", 409, False, "索引不存在"),
        ("VECTOR_DIMENSION_MISMATCH", 409, False, "不兼容"),
        ("VECTOR_SEARCH_FAILED", 502, True, "没有完成"),
    ]

    for code, status_code, retryable, message_fragment in cases:
        error = query_error_from_response(Response(code))
        assert error.code == code
        assert error.status_code == status_code
        assert error.retryable is retryable
        assert message_fragment in error.message


def test_vector_database_failure_is_visible_and_persisted_as_failed_answer(
    tmp_path,
    monkeypatch,
):
    class Response:
        is_success = False

        @staticmethod
        def json():
            return {
                "error": {
                    "code": "VECTOR_DB_UNAVAILABLE",
                    "message": "internal upstream message",
                    "retryable": True,
                }
            }

    class Client:
        def __init__(self, **_kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def post(self, url, json):
            if url.endswith("/route"):
                return _RouteResponse("knowledge")
            assert json
            return Response()

    monkeypatch.setattr(
        "frontend.product.services.conversations.httpx.AsyncClient",
        Client,
    )

    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "故障可见知识库", "description": ""},
        ).json()
        with client.product_session_factory() as session:
            session.add(
                Document(
                    knowledge_base_id=knowledge_base["id"],
                    display_name="failure-guide.pdf",
                    storage_path=str(tmp_path / "failure-guide.pdf"),
                    size_bytes=10,
                    sha256="f" * 64,
                    status="ready",
                )
            )
            session.get(KnowledgeBase, knowledge_base["id"]).index_manifest = json.dumps(
                {"schema_version": 1}
            )
            session.commit()
        conversation = client.post(
            "/api/conversations",
            json={"knowledge_base_id": knowledge_base["id"]},
        ).json()

        response = client.post(
            f"/api/conversations/{conversation['id']}/messages",
            json={"content": "这个问题不能被伪装成无答案"},
        )

        assert response.status_code == 503
        payload = response.json()
        request_id = payload["error"].pop("request_id")
        assert request_id
        assert payload == {
            "error": {
                "code": "VECTOR_DB_UNAVAILABLE",
                "message": "向量检索服务暂时不可用，请稍后重试。",
                "retryable": True,
            }
        }
        detail = client.get(f"/api/conversations/{conversation['id']}").json()
        assert len(detail["messages"]) == 2
        assistant = detail["messages"][1]
        assert assistant["status"] == "failed"
        assert assistant["answer_state"] == "failed"
        assert assistant["content"] == "向量检索服务暂时不可用，请稍后重试。"
        assert assistant["citations"] == []
        with client.product_session_factory() as session:
            run = session.scalar(
                select(AnswerRun).where(
                    AnswerRun.conversation_id == conversation["id"],
                )
            )
            assert run is not None
            assert run.status == "failed"
            assert run.request_id == request_id


def test_unexpected_retrieval_exception_still_persists_failed_answer_run(tmp_path, monkeypatch):
    class Client:
        def __init__(self, **_kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def post(self, url, json):
            if url.endswith("/route"):
                return _RouteResponse("knowledge")
            assert json["collection_names"]
            raise RuntimeError("retrieval exploded")

    monkeypatch.setattr(
        "frontend.product.services.conversations.httpx.AsyncClient",
        Client,
    )

    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "异常检索资料", "description": ""},
        ).json()
        with client.product_session_factory() as session:
            session.add(
                Document(
                    knowledge_base_id=knowledge_base["id"],
                    display_name="unexpected-failure.pdf",
                    storage_path=str(tmp_path / "unexpected-failure.pdf"),
                    size_bytes=10,
                    sha256="u" * 64,
                    status="ready",
                )
            )
            session.get(KnowledgeBase, knowledge_base["id"]).index_manifest = json.dumps({"schema_version": 1})
            session.commit()
        conversation = client.post(
            "/api/conversations",
            json={"knowledge_base_id": knowledge_base["id"]},
        ).json()

        response = client.post(
            f"/api/conversations/{conversation['id']}/messages",
            headers={"X-Request-ID": "unexpected-retrieval-1"},
            json={"content": "这个检索会直接抛异常"},
        )
        assert response.status_code == 503
        assert response.json()["error"]["request_id"] == "unexpected-retrieval-1"

        with client.product_session_factory() as session:
            run = session.scalar(
                select(AnswerRun).where(AnswerRun.conversation_id == conversation["id"])
            )
            assert run is not None
            assert run.status == "failed"
            assert run.request_id == "unexpected-retrieval-1"
            assert run.current_stage == "retrieval"
            assert run.failure_code == "QUERY_UNAVAILABLE"


def test_citation_maps_product_document_id_back_to_original_filename(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(
        "frontend.product.services.documents.UPLOAD_DIR",
        tmp_path / "uploads",
    )
    captured = {}

    class Response:
        is_success = True

        @staticmethod
        def json():
            return {
                "result": "Milvus 是向量数据库。",
                "trace": {
                    "iterations": [
                        {
                            "retrieved_documents": [
                                {
                                    "document_id": captured["document_id"],
                                    "display_name": f"{captured['document_id']}.pdf",
                                    "page_number": 1,
                                    "text": "What is Milvus?",
                                    "supported": True,
                                }
                            ]
                        }
                    ]
                },
            }

    class Client:
        def __init__(self, **_kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def post(self, url, json):
            if url.endswith("/route"):
                return _RouteResponse("knowledge")
            assert json
            return Response()

    monkeypatch.setattr(
        "frontend.product.services.conversations.httpx.AsyncClient",
        Client,
    )

    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "Milvus 资料", "description": ""},
        ).json()
        upload = client.post(
            f"/api/knowledge-bases/{knowledge_base['id']}/documents",
            files={
                "file": (
                    "WhatisMilvus.pdf",
                    b"%PDF-1.7\ncitation mapping",
                    "application/pdf",
                )
            },
        ).json()
        captured["document_id"] = upload["document"]["id"]
        with client.product_session_factory() as session:
            session.get(Document, captured["document_id"]).status = "ready"
            session.get(KnowledgeBase, knowledge_base["id"]).index_manifest = json.dumps(
                {"schema_version": 1}
            )
            session.commit()
        conversation = client.post(
            "/api/conversations",
            json={"knowledge_base_id": knowledge_base["id"]},
        ).json()

        answer = client.post(
            f"/api/conversations/{conversation['id']}/messages",
            json={"content": "Milvus 是什么？"},
        ).json()

    citation = answer["assistant_message"]["citations"][0]
    assert citation["document_id"] == captured["document_id"]
    assert citation["display_name"] == "WhatisMilvus.pdf"


def test_knowledge_health_trend_and_actions_api(tmp_path):
    with product_client(tmp_path) as client:
        created = client.post(
            "/api/knowledge-bases",
            json={"name": "健康测试", "description": ""},
        )
        assert created.status_code == 201
        knowledge_base_id = created.json()["id"]

        health = client.get(f"/api/knowledge-bases/{knowledge_base_id}/health")
        assert health.status_code == 200
        current = health.json()["current"]
        assert current["formula_version"] == "1.1"
        assert "level" in current
        assert "attribution" in current["metrics"]["retrieval"]

        snapshot_response = client.post(f"/api/knowledge-bases/{knowledge_base_id}/health/snapshot")
        assert snapshot_response.status_code == 201
        snapshot = snapshot_response.json()["snapshot"]
        assert snapshot["id"].startswith("khs_")
        assert "level" in snapshot

        trend = client.get(f"/api/knowledge-bases/{knowledge_base_id}/health/trend")
        assert trend.status_code == 200
        items = trend.json()["items"]
        assert len(items) == 1
        assert "level" in items[0]
        assert items[0]["overall"] is None  # empty knowledge base -> partial

        actions = client.post(
            f"/api/knowledge-bases/{knowledge_base_id}/health/actions/run",
            json={"actions": ["UPLOAD_DOCUMENTS"]},
        )
        assert actions.status_code == 200
        action_result = actions.json()
        assert action_result["results"][0]["status"] == "requires_user_action"
        assert action_result["snapshot"]["id"].startswith("khs_")
        assert action_result["delta"]["direction"] == "changed"

        bad = client.post(
            f"/api/knowledge-bases/{knowledge_base_id}/health/actions/run",
            json={"actions": ["NOPE"]},
        )
        assert bad.status_code == 400


def test_single_access_model_revocation_blocks_old_conversation(tmp_path):
    with product_client(tmp_path) as client:
        created = client.post(
            "/api/admin/users",
            json={
                "username": "member",
                "password": "password1234",
                "display_name": "Member",
                "role": "member",
            },
        )
        assert created.status_code == 201
        member_id = created.json()["user"]["id"]
        department = client.post(
            "/api/admin/departments",
            json={"name": "研发部"},
        )
        assert department.status_code == 201
        department_id = department.json()["department"]["id"]
        assigned = client.patch(
            f"/api/admin/users/{member_id}",
            json={"department_id": department_id},
        )
        assert assigned.status_code == 200
        with client.product_session_factory() as session:
            member_user = session.scalar(select(User).where(User.username == "member"))

        workspace = client.post(
            "/api/workspaces",
            json={"name": "team", "description": ""},
        )
        assert workspace.status_code == 201
        workspace_id = workspace.json()["id"]

        added = client.post(
            f"/api/workspaces/{workspace_id}/members",
            json={"username": "member", "role": "viewer"},
        )
        assert added.status_code == 201

        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "shared", "description": "", "workspace_id": workspace_id},
        )
        assert knowledge_base.status_code == 201
        knowledge_base_id = knowledge_base.json()["id"]

        access = client.put(
            f"/api/admin/departments/{department_id}/knowledge-access",
            json={"knowledge_base_ids": [knowledge_base_id]},
        )
        assert access.status_code == 200

        # member (viewer) can create a conversation while membership exists
        app.dependency_overrides[require_user] = lambda: member_user
        conversation = client.post(
            "/api/conversations",
            json={"knowledge_base_id": knowledge_base_id},
        )
        assert conversation.status_code == 201
        conversation_id = conversation.json()["id"]

        # Removing the department grant must invalidate the old fixed scope.
        app.dependency_overrides[require_user] = lambda: client.admin_user
        removed = client.put(
            f"/api/admin/departments/{department_id}/knowledge-access",
            json={"knowledge_base_ids": []},
        )
        assert removed.status_code == 200

        app.dependency_overrides[require_user] = lambda: member_user
        denied = client.post(
            f"/api/conversations/{conversation_id}/messages",
            json={"content": "hi"},
        )
        assert denied.status_code == 404  # resource hidden, no retriever touched


def test_team_trust_non_member_cannot_create_conversation(tmp_path):
    with product_client(tmp_path) as client:
        client.post(
            "/api/admin/users",
            json={
                "username": "stranger",
                "password": "password1234",
                "display_name": "Stranger",
                "role": "member",
            },
        )
        workspace = client.post(
            "/api/workspaces",
            json={"name": "private", "description": ""},
        ).json()
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={
                "name": "private-kb",
                "description": "",
                "workspace_id": workspace["id"],
            },
        ).json()

        with client.product_session_factory() as session:
            stranger = session.scalar(select(User).where(User.username == "stranger"))
        app.dependency_overrides[require_user] = lambda: stranger
        denied = client.post(
            "/api/conversations",
            json={"knowledge_base_id": knowledge_base["id"]},
        )
        assert denied.status_code == 404


def test_team_trust_viewer_cannot_write(tmp_path):
    with product_client(tmp_path) as client:
        client.post(
            "/api/admin/users",
            json={
                "username": "viewer",
                "password": "password1234",
                "display_name": "Viewer",
                "role": "member",
            },
        )
        workspace = client.post(
            "/api/workspaces",
            json={"name": "readonly", "description": ""},
        ).json()
        client.post(
            f"/api/workspaces/{workspace['id']}/members",
            json={"username": "viewer", "role": "viewer"},
        )
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={
                "name": "ro-kb",
                "description": "",
                "workspace_id": workspace["id"],
            },
        ).json()

        with client.product_session_factory() as session:
            viewer = session.scalar(select(User).where(User.username == "viewer"))
        app.dependency_overrides[require_user] = lambda: viewer
        denied = client.post(
            f"/api/knowledge-bases/{knowledge_base['id']}/reindex",
        )
        assert denied.status_code == 403


def test_kb_members_api_admin_only_and_flow(tmp_path):
    with product_client(tmp_path) as client:
        client.post(
            "/api/admin/users",
            json={
                "username": "viewer",
                "password": "password1234",
                "display_name": "Viewer",
                "role": "member",
            },
        )
        workspace = client.post(
            "/api/workspaces",
            json={"name": "team", "description": ""},
        ).json()
        client.post(
            f"/api/workspaces/{workspace['id']}/members",
            json={"username": "viewer", "role": "viewer"},
        )
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={
                "name": "kb",
                "description": "",
                "workspace_id": workspace["id"],
            },
        ).json()
        kb_id = knowledge_base["id"]

        with client.product_session_factory() as session:
            viewer = session.scalar(select(User).where(User.username == "viewer"))
        app.dependency_overrides[require_user] = lambda: viewer
        # viewer cannot read the member configuration (admin-only)
        assert client.get(f"/api/knowledge-bases/{kb_id}/members").status_code == 403
        assert (
            client.post(
                f"/api/knowledge-bases/{kb_id}/members",
                json={"username": "viewer", "role": "editor"},
            ).status_code
            == 403
        )

        # owner can manage KB overrides
        app.dependency_overrides[require_user] = lambda: client.admin_user
        added = client.post(
            f"/api/knowledge-bases/{kb_id}/members",
            json={"username": "viewer", "role": "editor"},
        )
        assert added.status_code == 201
        member_user_id = added.json()["member"]["user_id"]
        assert (
            client.get(f"/api/knowledge-bases/{kb_id}/members").json()["items"][0]["role"]
            == "editor"
        )
        patched = client.patch(
            f"/api/knowledge-bases/{kb_id}/members/{member_user_id}",
            json={"role": "viewer"},
        )
        assert patched.status_code == 200
        removed = client.delete(f"/api/knowledge-bases/{kb_id}/members/{member_user_id}")
        assert removed.status_code == 204


def test_groups_api_admin_only_and_flow(tmp_path):
    with product_client(tmp_path) as client:
        client.post(
            "/api/admin/users",
            json={
                "username": "member",
                "password": "password1234",
                "display_name": "Member",
                "role": "member",
            },
        )
        workspace = client.post(
            "/api/workspaces",
            json={"name": "team", "description": ""},
        ).json()
        workspace_id = workspace["id"]
        client.post(
            f"/api/workspaces/{workspace_id}/members",
            json={"username": "member", "role": "viewer"},
        )

        group = client.post(
            f"/api/workspaces/{workspace_id}/groups",
            json={"name": "editors", "role": "editor"},
        )
        assert group.status_code == 201
        group_id = group.json()["group"]["id"]
        added = client.post(
            f"/api/workspaces/{workspace_id}/groups/{group_id}/members",
            json={"username": "member"},
        )
        assert added.status_code == 201
        assert (
            client.get(f"/api/workspaces/{workspace_id}/groups/{group_id}").json()["members"][0][
                "username"
            ]
            == "member"
        )

        with client.product_session_factory() as session:
            member = session.scalar(select(User).where(User.username == "member"))
        app.dependency_overrides[require_user] = lambda: member
        assert client.get(f"/api/workspaces/{workspace_id}/groups").status_code == 403
        assert (
            client.post(
                f"/api/workspaces/{workspace_id}/groups",
                json={"name": "x", "role": "viewer"},
            ).status_code
            == 403
        )


# ---- P1-4.1 operation audit (ragent BizChangeLog equivalent) ----


def test_audit_logs_record_member_role_change_with_snapshots(tmp_path):
    with product_client(tmp_path) as client:
        created = client.post(
            "/api/admin/users",
            json={
                "username": "member",
                "password": "password1234",
                "display_name": "成员",
                "role": "member",
            },
        )
        assert created.status_code == 201
        workspace = client.post(
            "/api/workspaces",
            json={"name": "审计工作区", "description": ""},
        ).json()
        workspace_id = workspace["id"]
        added = client.post(
            f"/api/workspaces/{workspace_id}/members",
            json={"username": "member", "role": "viewer"},
        )
        assert added.status_code == 201
        member_user_id = added.json()["member"]["user_id"]

        changed = client.patch(
            f"/api/workspaces/{workspace_id}/members/{member_user_id}",
            json={"role": "editor"},
        )
        assert changed.status_code == 200

        page = client.get("/api/admin/audit-logs").json()
        assert page["total"] >= 2
        by_type = {item["operation_type"]: item for item in page["items"]}
        assert "ADD_WORKSPACE_MEMBER" in by_type
        assert "SET_MEMBER_ROLE" in by_type

        add = by_type["ADD_WORKSPACE_MEMBER"]
        assert add["biz_id"] == workspace_id
        assert add["before_snapshot"] is None
        assert add["after_snapshot"]["user_id"] == member_user_id
        assert add["after_snapshot"]["role"] == "viewer"

        set_role = by_type["SET_MEMBER_ROLE"]
        assert set_role["biz_id"] == workspace_id
        assert set_role["before_snapshot"]["role"] == "viewer"
        assert set_role["after_snapshot"]["role"] == "editor"
        assert set_role["success"] is True
        assert set_role["change_diff"]


def test_audit_logs_capture_operator_from_request_context(tmp_path):
    from frontend.product.services.audit import AuditContext, bind_audit_context

    with product_client(tmp_path) as client:
        bind_audit_context(
            AuditContext(
                operator_id=client.admin_user.id,
                operator_name="测试用户",
                operator_role="admin",
                ip="203.0.113.9",
            )
        )
        created = client.post(
            "/api/admin/users",
            json={
                "username": "context-user",
                "password": "password1234",
                "display_name": "上下文用户",
                "role": "member",
            },
        )
        assert created.status_code == 201
        page = client.get("/api/admin/audit-logs").json()
        create_user = next(
            item for item in page["items"] if item["operation_type"] == "CREATE_USER"
        )
        assert create_user["operator_id"] == client.admin_user.id
        assert create_user["operator_name"] == "测试用户"
        assert create_user["ip"] == "203.0.113.9"


def test_audit_logs_require_admin(tmp_path):
    with product_client(tmp_path) as client:
        client.post(
            "/api/admin/users",
            json={
                "username": "plain",
                "password": "password1234",
                "display_name": "普通用户",
                "role": "member",
            },
        )
        with client.product_session_factory() as session:
            plain = session.scalar(select(User).where(User.username == "plain"))
        app.dependency_overrides[require_user] = lambda: plain
        assert client.get("/api/admin/audit-logs").status_code == 403
        assert (
            client.get("/api/admin/audit-logs", params={"page": 1, "page_size": 1}).status_code
            == 403
        )


def test_audit_logs_pagination_and_filter(tmp_path):
    from frontend.product.services.audit import record_operation

    with product_client(tmp_path) as client:
        with client.product_session_factory() as session:
            for index in range(5):
                record_operation(
                    session,
                    biz_type="test_biz",
                    biz_id=f"biz-{index}",
                    operation_type="TEST_OP",
                    action_desc=f"测试操作 {index}",
                    before=None,
                    after={"index": index},
                    operator_id=client.admin_user.id,
                )

        page = client.get("/api/admin/audit-logs", params={"page": 1, "page_size": 2}).json()
        assert len(page["items"]) == 2
        assert page["total"] >= 5
        assert page["page"] == 1
        assert page["page_size"] == 2

        filtered = client.get(
            "/api/admin/audit-logs",
            params={
                "biz_type": "test_biz",
                "operation_type": "TEST_OP",
                "success": "true",
            },
        ).json()
        assert filtered["total"] == 5

        none = client.get("/api/admin/audit-logs", params={"operation_type": "NOPE"}).json()
        assert none["total"] == 0


# ---- P1-4.2 message feedback (ragent MessageFeedback equivalent) ----


def _make_feedback_context(client, name_suffix: str = ""):
    """Create a knowledge base, conversation and one completed assistant message."""
    knowledge_base = client.post(
        "/api/knowledge-bases",
        json={"name": f"反馈知识库{name_suffix}", "description": ""},
    ).json()
    conversation = client.post(
        "/api/conversations",
        json={"knowledge_base_id": knowledge_base["id"]},
    ).json()
    with client.product_session_factory() as session:
        assistant = Message(
            conversation_id=conversation["id"],
            role="assistant",
            content="回答内容",
            status="succeeded",
            answer_state="fully_grounded",
        )
        session.add(assistant)
        session.commit()
        message_id = assistant.id
    return knowledge_base, conversation, message_id


def test_message_feedback_submit_overwrite_cancel_and_restore(tmp_path):
    with product_client(tmp_path) as client:
        _, conversation, message_id = _make_feedback_context(client)
        feedback_url = f"/api/conversations/{conversation['id']}/messages/{message_id}/feedback"

        submitted = client.post(feedback_url, json={"vote": 1})
        assert submitted.status_code == 200
        assert submitted.json()["feedback"] == {"vote": 1, "cancelled": False}
        with client.product_session_factory() as session:
            rows = session.scalars(select(MessageFeedback)).all()
            assert len(rows) == 1
            assert rows[0].vote == 1
            assert rows[0].cancelled is False

        overwritten = client.post(
            feedback_url,
            json={"vote": -1, "reason": "回答不完整"},
        )
        assert overwritten.status_code == 200
        assert overwritten.json()["feedback"] == {"vote": -1, "cancelled": False}
        with client.product_session_factory() as session:
            rows = session.scalars(select(MessageFeedback)).all()
            assert len(rows) == 1  # idempotent overwrite, not a second row
            assert rows[0].vote == -1
            assert rows[0].reason == "回答不完整"

        detail = client.get(f"/api/conversations/{conversation['id']}").json()
        assistant = next(message for message in detail["messages"] if message["id"] == message_id)
        assert assistant["feedback"] == {"vote": -1, "cancelled": False}

        cancelled = client.delete(feedback_url)
        assert cancelled.status_code == 204
        with client.product_session_factory() as session:
            rows = session.scalars(select(MessageFeedback)).all()
            assert len(rows) == 1
            assert rows[0].vote is None
            assert rows[0].cancelled is True
            assert rows[0].reason is None

        restored = client.post(feedback_url, json={"vote": 1})
        assert restored.status_code == 200
        assert restored.json()["feedback"] == {"vote": 1, "cancelled": False}


def test_message_feedback_cancel_is_idempotent_without_prior_vote(tmp_path):
    with product_client(tmp_path) as client:
        _, conversation, message_id = _make_feedback_context(client)
        feedback_url = f"/api/conversations/{conversation['id']}/messages/{message_id}/feedback"

        assert client.delete(feedback_url).status_code == 204
        assert client.delete(feedback_url).status_code == 204
        with client.product_session_factory() as session:
            rows = session.scalars(select(MessageFeedback)).all()
            assert len(rows) == 1
            assert rows[0].cancelled is True
            assert rows[0].vote is None


def test_message_feedback_target_and_ownership_errors(tmp_path):
    with product_client(tmp_path) as client:
        _, conversation_a, message_a = _make_feedback_context(client, name_suffix="-A")
        _, conversation_b, message_b = _make_feedback_context(client, name_suffix="-B")
        url_a = f"/api/conversations/{conversation_a['id']}/messages/{message_a}/feedback"

        # Message from another conversation -> 404 without leaking existence.
        wrong = client.post(
            f"/api/conversations/{conversation_a['id']}/messages/{message_b}/feedback",
            json={"vote": 1},
        )
        assert wrong.status_code == 404

        # Missing message id inside a real conversation -> 404.
        missing = client.post(
            f"/api/conversations/{conversation_a['id']}/messages/nope_123/feedback",
            json={"vote": 1},
        )
        assert missing.status_code == 404

        # Missing conversation -> 404.
        ghost = client.post(
            "/api/conversations/ghost_conv/messages/x/feedback",
            json={"vote": 1},
        )
        assert ghost.status_code == 404

        # User messages are not feedback targets -> 400.
        with client.product_session_factory() as session:
            user_message = Message(
                conversation_id=conversation_a["id"],
                role="user",
                content="提问",
                status="succeeded",
            )
            session.add(user_message)
            session.commit()
            user_message_id = user_message.id
        user_target = client.post(
            f"/api/conversations/{conversation_a['id']}/messages/{user_message_id}/feedback",
            json={"vote": 1},
        )
        assert user_target.status_code == 400

        # Invalid votes rejected by the schema -> 422.
        for bad_vote in (0, 2):
            assert client.post(url_a, json={"vote": bad_vote}).status_code == 422

        # A conversation owned by another user -> 404.
        with client.product_session_factory() as session:
            stranger = User(
                username="stranger",
                display_name="陌生人",
                password_hash="x" * 60,
                role="member",
            )
            session.add(stranger)
            session.commit()
            stranger_id = stranger.id
        with client.product_session_factory() as session:
            stranger = session.scalar(select(User).where(User.id == stranger_id))
        app.dependency_overrides[require_user] = lambda: stranger
        try:
            foreign = client.post(url_a, json={"vote": 1})
            assert foreign.status_code == 404
        finally:
            app.dependency_overrides.pop(require_user, None)


def test_message_feedback_negative_rate_enters_health_snapshot(tmp_path):
    with product_client(tmp_path) as client:
        knowledge_base, conversation, message_id = _make_feedback_context(client)
        with client.product_session_factory() as session:
            for _ in range(3):
                session.add(
                    Message(
                        conversation_id=conversation["id"],
                        role="assistant",
                        content="补充回答",
                        status="succeeded",
                        answer_state="fully_grounded",
                    )
                )
            session.commit()

        negative = client.post(
            f"/api/conversations/{conversation['id']}/messages/{message_id}/feedback",
            json={"vote": -1},
        )
        assert negative.status_code == 200

        snapshot = client.post(f"/api/knowledge-bases/{knowledge_base['id']}/health/snapshot")
        assert snapshot.status_code == 201
        retrieval = snapshot.json()["snapshot"]["metrics"]["retrieval"]
        assert retrieval["message_sample_count"] == 4
        assert retrieval["negative_feedback_rate"] == 0.25


# ---- P4 admin dashboard ----


def test_admin_dashboard_overview_admin_ok_and_viewer_forbidden(tmp_path):
    with product_client(tmp_path) as client:
        assert client.get("/api/admin/dashboard/overview").status_code == 200
        assert client.get("/api/admin/dashboard/trends?days=7").status_code == 200

        client.post(
            "/api/admin/users",
            json={
                "username": "dash-viewer",
                "password": "password1234",
                "display_name": "查看者",
                "role": "member",
            },
        )
        with client.product_session_factory() as session:
            viewer = session.scalar(select(User).where(User.username == "dash-viewer"))
        app.dependency_overrides[require_user] = lambda: viewer
        try:
            assert client.get("/api/admin/dashboard/overview").status_code == 403
            assert client.get("/api/admin/dashboard/trends?days=7").status_code == 403
        finally:
            app.dependency_overrides.pop(require_user, None)


def test_admin_dashboard_trends_days_validation(tmp_path):
    with product_client(tmp_path) as client:
        assert client.get("/api/admin/dashboard/trends?days=7").status_code == 200
        assert client.get("/api/admin/dashboard/trends?days=30").status_code == 200
        for bad in (6, 14, 31):
            assert (
                client.get(f"/api/admin/dashboard/trends?days={bad}").status_code == 422
            )


def test_admin_knowledge_inventory_and_ingestion_routes_are_admin_only(tmp_path):
    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "管理员资料库", "description": "运营检查"},
        ).json()

        inventory = client.get("/api/admin/knowledge-bases")
        assert inventory.status_code == 200
        assert [item["id"] for item in inventory.json()["items"]] == [knowledge_base["id"]]

        documents = client.get(
            f"/api/admin/knowledge-bases/{knowledge_base['id']}/documents"
        )
        assert documents.status_code == 200
        assert documents.json() == {"items": []}

        health = client.get(
            f"/api/admin/knowledge-bases/{knowledge_base['id']}/health"
        )
        assert health.status_code == 200
        assert health.json()["snapshot"] is None
        assert health.json()["current"]["status"] == "partial"

        jobs = client.get("/api/admin/ingest-jobs")
        assert jobs.status_code == 200
        assert jobs.json()["total"] == 0

        client.post(
            "/api/admin/users",
            json={
                "username": "inventory-member",
                "password": "password1234",
                "display_name": "普通成员",
                "role": "member",
            },
        )
        with client.product_session_factory() as session:
            viewer = session.scalar(select(User).where(User.username == "inventory-member"))
        app.dependency_overrides[require_user] = lambda: viewer
        try:
            assert client.get("/api/admin/knowledge-bases").status_code == 403
            assert client.get(
                f"/api/admin/knowledge-bases/{knowledge_base['id']}/documents"
            ).status_code == 403
            assert client.get(
                f"/api/admin/knowledge-bases/{knowledge_base['id']}/health"
            ).status_code == 403
            assert client.get("/api/admin/ingest-jobs").status_code == 403
        finally:
            app.dependency_overrides.pop(require_user, None)


def test_admin_knowledge_inventory_includes_latest_health_summary(tmp_path):
    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "带健康快照的资料库", "description": ""},
        ).json()
        with client.product_session_factory() as session:
            session.add(
                KnowledgeHealthSnapshot(
                    knowledge_base_id=knowledge_base["id"],
                    owner_id=client.admin_user.id,
                    formula_version="1.1",
                    status="complete",
                    overall_score=72.5,
                    data_score=80,
                    retrieval_score=70,
                    trust_score=65,
                    metrics={},
                    deductions=[],
                    actions=[],
                )
            )
            session.commit()

        inventory = client.get("/api/admin/knowledge-bases")

    assert inventory.status_code == 200
    item = inventory.json()["items"][0]
    assert item["health_status"] == "healthy"
    assert item["health_score"] == 72.5
    assert item["health_updated_at"]


def test_admin_runs_expose_only_persisted_answer_evidence_and_are_admin_only(tmp_path):
    with product_client(tmp_path) as client:
        knowledge_base = client.post(
            "/api/knowledge-bases",
            json={"name": "Trace 资料库", "description": "真实运行记录"},
        ).json()
        conversation = client.post(
            "/api/conversations",
            json={"knowledge_base_id": knowledge_base["id"]},
        ).json()

        with client.product_session_factory() as session:
            session.add_all(
                [
                    Message(
                        conversation_id=conversation["id"],
                        role="user",
                        content="这个回答的依据是什么？",
                        status="succeeded",
                    ),
                    Message(
                        conversation_id=conversation["id"],
                        role="assistant",
                        content="这是一个真实持久化的回答。",
                        status="succeeded",
                        answer_state="fully_grounded",
                        trust_status="fully_grounded",
                        policy_action="allow",
                        risk_level="low",
                    ),
                ]
            )
            session.commit()
            assistant = session.scalar(
                select(Message).where(
                    Message.conversation_id == conversation["id"],
                    Message.role == "assistant",
                )
            )
            assistant_id = assistant.id
            session.add(
                MessageFeedback(
                    message_id=assistant_id,
                    user_id=client.admin_user.id,
                    vote=-1,
                    reason="证据不足",
                    comment="希望看到更明确的来源。",
                )
            )
            session.commit()

        response = client.get(
            "/api/admin/runs?trust_status=fully_grounded&risk_level=low&feedback=negative"
        )
        assert response.status_code == 200
        payload = response.json()
        assert payload["total"] == 1
        item = payload["items"][0]
        assert item["message"]["id"] == assistant_id
        assert item["message"]["trust_status"] == "fully_grounded"
        assert item["message"]["citations"] == []
        assert item["message"]["claims"] == []
        assert item["knowledge_base"]["id"] == knowledge_base["id"]
        assert item["feedback"] == {"positive": 0, "negative": 1, "comment_count": 1}
        assert len(item["feedback_items"]) == 1
        assert item["feedback_items"][0]["vote"] == -1
        assert item["feedback_items"][0]["reason"] == "证据不足"
        assert item["feedback_items"][0]["comment"] == "希望看到更明确的来源。"
        assert item["feedback_items"][0]["created_at"]

        detail = client.get(f"/api/admin/runs/{assistant_id}")
        assert detail.status_code == 200
        assert detail.json()["message"]["content"] == "这是一个真实持久化的回答。"

        client.post(
            "/api/admin/users",
            json={
                "username": "trace-member",
                "password": "password1234",
                "display_name": "Trace 成员",
                "role": "member",
            },
        )
        with client.product_session_factory() as session:
            viewer = session.scalar(select(User).where(User.username == "trace-member"))
        app.dependency_overrides[require_user] = lambda: viewer
        try:
            assert client.get("/api/admin/runs").status_code == 403
            assert client.get(f"/api/admin/runs/{assistant_id}").status_code == 403
        finally:
            app.dependency_overrides.pop(require_user, None)


def test_admin_runs_filter_auto_scope_by_execution_snapshot(tmp_path):
    with product_client(tmp_path) as client:
        first = client.post("/api/knowledge-bases", json={"name": "范围 A"}).json()
        second = client.post("/api/knowledge-bases", json={"name": "范围 B"}).json()
        conversation = client.post("/api/conversations", json={}).json()

        with client.product_session_factory() as session:
            question = Message(
                conversation_id=conversation["id"],
                role="user",
                content="自动范围问题",
                status="succeeded",
            )
            answer = Message(
                conversation_id=conversation["id"],
                role="assistant",
                content="自动范围回答",
                status="succeeded",
            )
            session.add_all([question, answer])
            session.flush()
            session.add(
                AnswerRun(
                    conversation_id=conversation["id"],
                    question_message_id=question.id,
                    answer_message_id=answer.id,
                    status="succeeded",
                    total_latency_ms=120,
                    query_scope_snapshot={
                        "schema_version": 1,
                        "mode": "auto",
                        "knowledge_base_ids": [first["id"], second["id"]],
                        "collection_names": ["collection_a", "collection_b"],
                        "resolved_at": "2026-08-20T10:00:00+08:00",
                        "status_counts": {
                            "usable": 2,
                            "no_documents": 0,
                            "processing": 0,
                            "failed": 0,
                            "needs_rebuild": 0,
                            "unknown": 0,
                        },
                    },
                )
            )
            session.commit()

        response = client.get(f"/api/admin/runs?knowledge_base_id={second['id']}")
        assert response.status_code == 200
        assert response.json()["total"] == 1
        assert response.json()["items"][0]["scope"]["knowledge_base_ids"] == [first["id"], second["id"]]


# A minimal valid 1x1 transparent PNG (image upload without OCR).
PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def test_query_scope_reports_partial_counts_and_auto_conversation_has_no_kb_anchor(tmp_path):
    with product_client(tmp_path) as client:
        first = client.post("/api/knowledge-bases", json={"name": "可用资料"}).json()
        second = client.post("/api/knowledge-bases", json={"name": "准备中资料"}).json()
        with client.product_session_factory() as session:
            session.add(
                Document(
                    knowledge_base_id=first["id"],
                    display_name="ready.pdf",
                    storage_path=str(tmp_path / "ready.pdf"),
                    size_bytes=1,
                    sha256="r" * 64,
                    status="ready",
                )
            )
            session.add(
                Document(
                    knowledge_base_id=second["id"],
                    display_name="processing.pdf",
                    storage_path=str(tmp_path / "processing.pdf"),
                    size_bytes=1,
                    sha256="p" * 64,
                    status="processing",
                )
            )
            session.get(KnowledgeBase, first["id"]).index_manifest = json.dumps({"schema_version": 1})
            session.commit()

        scope = client.get("/api/me/query-scope")
        assert scope.status_code == 200
        assert scope.json()["state"] == "partial"
        assert scope.json()["askable"] is True
        assert scope.json()["primary_action"] is None
        counts = scope.json()["status_counts"]
        assert counts == {
            "usable": 1,
            "no_documents": 0,
            "processing": 1,
            "failed": 0,
            "needs_rebuild": 0,
            "unknown": 0,
        }
        assert sum(counts.values()) == scope.json()["accessible_knowledge_base_count"]

        conversation = client.post("/api/conversations", json={}).json()
        assert conversation["scope_mode"] == "auto"
        assert conversation["knowledge_base_id"] is None


def test_fixed_scope_is_reauthorized_after_department_access_revoke(tmp_path):
    with product_client(tmp_path) as client:
        knowledge_base = client.post("/api/knowledge-bases", json={"name": "撤权验证"}).json()
        with client.product_session_factory() as session:
            member = User(
                username="scope-member",
                display_name="Scope 成员",
                password_hash="disabled",
                role="member",
            )
            session.add(member)
            session.flush()
            department = Department(name="Scope 部门")
            session.add(department)
            session.flush()
            member.department_id = department.id
            knowledge_base_row = session.get(KnowledgeBase, knowledge_base["id"])
            session.add(
                Document(
                    knowledge_base_id=knowledge_base["id"],
                    display_name="scope.pdf",
                    storage_path=str(tmp_path / "scope.pdf"),
                    size_bytes=1,
                    sha256="q" * 64,
                    status="ready",
                )
            )
            knowledge_base_row.index_manifest = json.dumps({"schema_version": 1})
            department_access = DepartmentKnowledgeBase(
                department_id=department.id,
                knowledge_base_id=knowledge_base_row.id,
            )
            session.add(department_access)
            conversation = Conversation(
                owner_id=member.id,
                knowledge_base_id=knowledge_base["id"],
                scope_mode="fixed",
            )
            session.add(conversation)
            session.commit()
            session.refresh(conversation)
            assert resolve_conversation_scope(session, member, conversation).payload["askable"] is True
            session.delete(department_access)
            session.commit()
            with pytest.raises(ProductError) as exc_info:
                resolve_conversation_scope(session, member, conversation)
            assert exc_info.value.status_code == 404


def test_answer_run_persists_execution_scope_snapshot(tmp_path, monkeypatch):
    class Response:
        is_success = True

        @staticmethod
        def json():
            return {"result": "可审计回答", "provider": "test", "model": "test-model", "trace": {}}

    class Client:
        def __init__(self, **_kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def post(self, url, json):
            if url.endswith("/route"):
                return _RouteResponse("knowledge")
            assert json["collection_names"]
            return Response()

    monkeypatch.setattr("frontend.product.services.conversations.httpx.AsyncClient", Client)

    with product_client(tmp_path) as client:
        knowledge_base = client.post("/api/knowledge-bases", json={"name": "快照资料"}).json()
        with client.product_session_factory() as session:
            session.add(
                Document(
                    knowledge_base_id=knowledge_base["id"],
                    display_name="snapshot.pdf",
                    storage_path=str(tmp_path / "snapshot.pdf"),
                    size_bytes=1,
                    sha256="t" * 64,
                    status="ready",
                )
            )
            session.get(KnowledgeBase, knowledge_base["id"]).index_manifest = json.dumps({"schema_version": 1})
            session.commit()
        conversation = client.post("/api/conversations", json={}).json()
        response = client.post(
            f"/api/conversations/{conversation['id']}/messages",
            json={"content": "快照问题"},
        )
        assert response.status_code == 200
        with client.product_session_factory() as session:
            run = session.scalar(select(AnswerRun).where(AnswerRun.conversation_id == conversation["id"]))
            assert run is not None
            assert run.status == "succeeded"
            assert run.answer_message_id is not None
            assert run.query_scope_snapshot["schema_version"] == 1
            assert run.query_scope_snapshot["mode"] == "auto"
            assert run.query_scope_snapshot["knowledge_base_ids"] == [knowledge_base["id"]]
            assert run.query_scope_snapshot["collection_names"]
            assert run.provider == "test"
            assert run.model == "test-model"
