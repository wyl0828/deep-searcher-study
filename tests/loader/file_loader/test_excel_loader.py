"""ExcelLoader tests (aligned with ragent ExcelDocumentParser/TableChunker)."""

from __future__ import annotations

import io
from pathlib import Path

from deepsearcher.loader.file_loader.excel_loader import ExcelLoader

FIXTURES = Path(__file__).parent / "fixtures"


def _openpyxl_bytes(build) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    build(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_formula_cached_fixture_reads_cache_and_merges(tmp_path):
    fixture = FIXTURES / "formula-cached.xlsx"
    documents = ExcelLoader().load_file(str(fixture))
    assert len(documents) == 1
    document = documents[0]
    assert document.metadata["kind"] == "table"
    assert document.metadata["sheet_name"] == "数据"
    assert document.metadata["sheet_count"] == 1
    assert document.metadata["_table_headers"] == ["项目", "金额"]

    rows = document.metadata["_table_rows"]
    # header rows: 项目/金额 ; 数据行: A/100, B/200, 合计/300, 合并单元格展开
    assert ["A", "100"] in rows
    assert ["B", "200"] in rows
    assert ["合计", "300"] in rows  # cached formula value 300
    merged = [row for row in rows if row[0] == "合并单元格"]
    assert merged and merged[0][1] == "合并单元格"  # top-left value expanded


def test_formula_without_cache_falls_back_to_formula_string(tmp_path):
    def build(workbook):
        sheet = workbook.active
        sheet.title = "S"
        sheet.append(["a", "b"])
        sheet.append([1, 2])
        sheet.append(["sum", "=B2+B3"])

    path = tmp_path / "formula.xlsx"
    path.write_bytes(_openpyxl_bytes(build))
    documents = ExcelLoader().load_file(str(path))
    rows = documents[0].metadata["_table_rows"]
    assert ["sum", "=B2+B3"] in rows


def test_hidden_sheet_skipped_and_sheet_count_visible(tmp_path):
    def build(workbook):
        visible = workbook.active
        visible.title = "可见"
        visible.append(["x", "y"])
        workbook.create_sheet("隐藏")
        workbook["隐藏"].sheet_state = "hidden"
        workbook.create_sheet("可见2").append(["p", "q"])

    path = tmp_path / "hidden.xlsx"
    path.write_bytes(_openpyxl_bytes(build))
    documents = ExcelLoader().load_file(str(path))
    assert len(documents) == 2
    assert documents[0].metadata["sheet_count"] == 2
    names = {document.metadata["sheet_name"] for document in documents}
    assert names == {"可见", "可见2"}


def test_multi_header_rows_flatten(tmp_path):
    def build(workbook):
        sheet = workbook.active
        sheet.title = "多行表头"
        sheet.append(["财务", "财务"])
        sheet.append(["收入", "支出"])
        sheet.append([100, 50])

    path = tmp_path / "headers.xlsx"
    path.write_bytes(_openpyxl_bytes(build))
    documents = ExcelLoader(header_rows=2).load_file(str(path))
    headers = documents[0].metadata["_table_headers"]
    assert headers == ["财务|收入", "财务|支出"]


def test_cell_pipe_and_newline_escaped(tmp_path):
    def build(workbook):
        sheet = workbook.active
        sheet.title = "转义"
        sheet.append(["名称", "备注"])
        sheet.append(["a|b", "line1\nline2"])

    path = tmp_path / "escape.xlsx"
    path.write_bytes(_openpyxl_bytes(build))
    documents = ExcelLoader().load_file(str(path))
    text = documents[0].page_content
    assert "a\\|b" in text
    assert "<br>" in text
