"""PptxLoader tests (slide-level multi-block documents)."""

from __future__ import annotations

import io

from deepsearcher.loader.file_loader.pptx_loader import PptxLoader


def _pptx_bytes() -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "季度汇报"
    text_box = slide.shapes.add_textbox(0, 100, 200, 50)
    text_box.text_frame.text = "正文内容"
    table_shape = slide.shapes.add_table(3, 2, 0, 200, 300, 100)
    table = table_shape.table
    for row_index, row in enumerate([("指标", "数值"), ("收入", "100"), ("支出", "50")]):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = value
    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "第二页"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_slide_level_text_and_table_blocks(tmp_path):
    path = tmp_path / "deck.pptx"
    path.write_bytes(_pptx_bytes())
    documents = PptxLoader().load_file(str(path))

    kinds = [document.metadata["kind"] for document in documents]
    assert "text" in kinds
    assert "table" in kinds
    for document in documents:
        assert document.metadata["total_slides"] == 2
        assert document.metadata["slide_number"] in (1, 2)
        assert document.metadata["extraction_method"] == "pptx"

    table = next(document for document in documents if document.metadata["kind"] == "table")
    assert table.metadata["slide_number"] == 1
    assert table.metadata["_table_headers"] == ["指标", "数值"]
    assert table.metadata["_table_rows"] == [["收入", "100"], ["支出", "50"]]
    assert "| 指标 | 数值 |" in table.page_content
